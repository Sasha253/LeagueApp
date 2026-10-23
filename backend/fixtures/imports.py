"""Importing a Season's Fixtures from the Admin's spreadsheet (FR-53), like chunk 4's import.

One row per Fixture: Date, Time (South African, 24-hour), Night (optional), Division, Pitch,
Home, Away and Friendly eligible (optional). The Season is chosen on the screen, not in the
file. A row is the same Fixture as one the club has when its Division, Home, Away and
kick-off match, so importing the same file twice creates nothing twice (docs/DECISIONS.md):

- a new Fixture is created, scheduled;
- an existing Fixture takes the file's Friendly eligible, if the cell isn't blank; nothing
  else of it changes, since moving a Fixture is a reschedule (chunk 16), not an import;
- Fixtures that aren't in the file are left alone.

plan() checks every row and writes nothing. apply() writes a plan with no errors.
"""

import csv
import datetime as dt
import io
import re
from dataclasses import dataclass, field

from django.db import models
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font

from clubs.imports import parse_weekday
from clubs.models import Club, Division, Field, Night, Season
from core.tables import Row
from teams.models import Team

from . import lifecycle
from .models import Fixture, FixtureTransition

COLUMNS = {
    "date": ("Date", "Match date"),
    "time": ("Time", "Kick-off", "Kick-off time"),
    "night": ("Night", "Weekday", "Day"),
    "division": ("Division", "Division name"),
    "pitch": ("Pitch", "Field"),
    "home": ("Home", "Home team"),
    "away": ("Away", "Away team"),
    "friendly_eligible": ("Friendly eligible", "Friendly"),
}
REQUIRED = {"date", "time", "division", "pitch", "home", "away"}

TEMPLATE_ROWS = [
    ("2026-10-27", "19:40", "Tuesday", "Division 2", "Pitch 3", "Red Lions", "Old Boys", "yes"),
    ("2026-10-27", "20:30", "Tuesday", "Division 2", "Pitch 3", "City Rovers", "Hat Trick FC", ""),
]
TEMPLATE_HELP = [
    ("Column", "What to put in it"),
    ("Date", "The match date: 2026-10-27, 27/10/2026 or 27 Oct 2026. Required."),
    ("Time", "Kick-off in South African time, 24-hour: 19:40. Required."),
    ("Night", "The weekday, e.g. Tuesday. Blank: the date's weekday. Must match the date."),
    ("Division", "The Division's name, e.g. Division 2. Required."),
    ("Pitch", "The pitch's name as in Club setup, e.g. Pitch 3. Required."),
    ("Home", "A confirmed Team of that Division this Season. Required."),
    ("Away", "Another confirmed Team of that Division this Season. Required."),
    ("Friendly eligible", "yes or no: may it become a friendly if it can't go ahead. Blank: no."),
    ("", "Every row is for the Season chosen in the app."),
    ("", "Re-importing the same file creates nothing twice."),
    ("", "A row is the same Fixture when Division, Home, Away, date and time match."),
    ("", "Fixtures missing from the file are left as they are. The import never moves one."),
]

MONTHS = {
    name: number
    for number, name in enumerate(
        ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"],
        start=1,
    )
}
ISO_DATE = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T]00:00(?::00)?)?$")
DMY_DATE = re.compile(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})$")
WORD_DATE = re.compile(r"^(?:([A-Za-z]{3,9}),?\s+)?(\d{1,2})\s+([A-Za-z]{3,9})\.?,?\s+(\d{4})$")
TIME = re.compile(r"^(\d{1,2})[:hH.](\d{2})(?::00)?$")
# Excel can keep a time as a moment on its day zero; only the time of day counts then.
EXCEL_TIME = re.compile(r"^1899-12-3[01] (\d{1,2}:\d{2}(?::00)?)$")
YES = {"yes", "y", "true", "1"}
NO = {"no", "n", "false", "0"}
DATE_HELP = "like 2026-10-27, 27/10/2026 or 27 Oct 2026"


class Action(models.TextChoices):
    CREATE = "create", "New Fixture"
    UPDATE = "update", "Fixture changes"
    UNCHANGED = "unchanged", "No change"
    ERROR = "error", "Has errors"


@dataclass
class PlannedRow:
    row: int
    kick_off: dt.datetime | None = None  # Aware.
    night: int | None = None
    division: str = ""
    pitch: str = ""
    home: str = ""
    away: str = ""
    friendly_eligible: bool | None = None  # What the Fixture will have, after the import.
    action: str = Action.ERROR
    changes: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    existing: Fixture | None = None
    division_obj: Division | None = None
    pitch_obj: Field | None = None
    home_obj: Team | None = None
    away_obj: Team | None = None


@dataclass
class Plan:
    rows: list[PlannedRow]

    @property
    def has_errors(self) -> bool:
        return any(row.errors for row in self.rows)

    @property
    def summary(self) -> dict[str, int]:
        def count(action):
            return sum(1 for row in self.rows if row.action == action)

        return {
            "rows": len(self.rows),
            "fixtures_created": count(Action.CREATE),
            "fixtures_updated": count(Action.UPDATE),
            "unchanged": count(Action.UNCHANGED),
            "errors": count(Action.ERROR),
        }


def parse_date(text: str) -> dt.date:
    plain = text.strip()
    weekday_word = None
    if match := ISO_DATE.fullmatch(plain):
        year, month, day = (int(part) for part in match.groups())
    elif match := DMY_DATE.fullmatch(plain):
        day, month, year = (int(part) for part in match.groups())
    elif match := WORD_DATE.fullmatch(plain):
        weekday_word, day_text, month_word, year_text = match.groups()
        month = MONTHS.get(month_word[:3].lower(), 0)
        if not month or not _month_name_fits(month_word, month):
            raise ValueError(f"Date must be a date, {DATE_HELP}.")
        day, year = int(day_text), int(year_text)
    else:
        raise ValueError(f"Date must be a date, {DATE_HELP}.")
    try:
        date = dt.date(year, month, day)
    except ValueError as error:
        raise ValueError(f"Date {plain} doesn't exist.") from error
    if weekday_word is not None and parse_weekday(weekday_word) != date.weekday():
        raise ValueError(f"{plain} isn't a {weekday_word}.")
    return date


def _month_name_fits(word: str, month: int) -> bool:
    full = dt.date(2000, month, 1).strftime("%B").lower()
    return full.startswith(word.lower().rstrip("."))


def parse_time(text: str) -> dt.time:
    plain = text.strip()
    if excel := EXCEL_TIME.fullmatch(plain):
        plain = excel.group(1)
    match = TIME.fullmatch(plain)
    if match:
        hour, minute = int(match.group(1)), int(match.group(2))
        if hour < 24 and minute < 60:
            return dt.time(hour, minute)
    raise ValueError("Time must be a kick-off time, 24-hour, like 19:40.")


def parse_flag(text: str) -> bool | None:
    word = text.strip().lower()
    if not word:
        return None
    if word in YES:
        return True
    if word in NO:
        return False
    raise ValueError("Friendly eligible must be yes, no or blank.")


def local_kick_off(date: dt.date, time: dt.time) -> dt.datetime:
    """The UTC moment of a South African date and time (the club's TIME_ZONE)."""
    local = timezone.make_aware(dt.datetime.combine(date, time), timezone.get_default_timezone())
    return local.astimezone(dt.UTC)


def _by_name(records, name_of) -> tuple[dict, dict]:
    exact: dict = {}
    loose: dict = {}
    for record in records:
        key, name = name_of(record)
        exact[(key, name)] = record
        loose.setdefault((key, name.casefold()), []).append(record)
    return exact, loose


def _find(exact: dict, loose: dict, key, name: str):
    """The record with exactly this name, else the only one matching it ignoring case.

    Raises LookupError when several match ignoring case.
    """
    found = exact.get((key, name))
    if found is not None:
        return found
    matches = loose.get((key, name.casefold()), [])
    if len(matches) > 1:
        raise LookupError
    return matches[0] if matches else None


def _day(date: dt.date) -> str:
    """27 Oct 2026"""
    return f"{date.day} {date:%b %Y}"


def _weekday_name(weekday: int) -> str:
    return Night.Weekday(weekday).label


def plan(club: Club, season: Season, rows: list[Row]) -> Plan:
    divisions = _by_name(
        Division.objects.filter(night__club=club).select_related("night"),
        lambda d: (d.night.weekday, d.name.strip()),
    )
    pitches = _by_name(Field.objects.filter(club=club), lambda f: (None, f.name.strip()))
    # Team names are unique per Division and Season ignoring case, among Teams not declined.
    teams = {
        (team.division_id, team.name.strip().casefold()): team
        for team in Team.objects.filter(season=season).exclude(status=Team.Status.DECLINED)
    }
    existing = list(
        Fixture.objects.filter(season=season).select_related("pitch", "home_team", "away_team")
    )
    by_key = {(f.division_id, f.home_team_id, f.away_team_id, f.kick_off): f for f in existing}
    # Who and what is busy at each kick-off: the club's Fixtures that aren't cancelled.
    busy = Fixture.objects.filter(club=club).exclude(status=Fixture.Status.CANCELLED)
    pitch_busy = {(f.pitch_id, f.kick_off): f for f in busy}
    team_busy: dict[tuple[int, dt.datetime], Fixture] = {}
    for fixture in busy:
        team_busy[(fixture.home_team_id, fixture.kick_off)] = fixture
        team_busy[(fixture.away_team_id, fixture.kick_off)] = fixture

    now = timezone.now()
    seen_fixture: dict[tuple, int] = {}
    seen_pitch: dict[tuple, int] = {}
    seen_team: dict[tuple, int] = {}
    planned = []

    for row in rows:
        result = PlannedRow(row=row.number)
        planned.append(result)
        cells = row.cells
        errors = result.errors

        date = time = flag = None
        if not cells.get("date"):
            errors.append("Date is empty.")
        else:
            try:
                date = parse_date(cells["date"])
            except ValueError as error:
                errors.append(str(error))
        if not cells.get("time"):
            errors.append("Time is empty.")
        else:
            try:
                time = parse_time(cells["time"])
            except ValueError as error:
                errors.append(str(error))
        try:
            flag = parse_flag(cells.get("friendly_eligible", ""))
        except ValueError as error:
            errors.append(str(error))

        night_text = cells.get("night", "")
        if night_text:
            result.night = parse_weekday(night_text)
            if result.night is None:
                errors.append(f'Night "{night_text}" isn\'t a weekday.')
            elif date is not None and date.weekday() != result.night:
                errors.append(
                    f"Night is {_weekday_name(result.night)}, but {_day(date)} is a "
                    f"{_weekday_name(date.weekday())}."
                )
        elif date is not None:
            result.night = date.weekday()

        for key, label in [("division", "Division"), ("pitch", "Pitch")]:
            setattr(result, key, cells.get(key, ""))
            if not getattr(result, key):
                errors.append(f"{label} is empty.")
        result.home, result.away = cells.get("home", ""), cells.get("away", "")
        if not result.home:
            errors.append("Home is empty.")
        if not result.away:
            errors.append("Away is empty.")
        if result.home and result.home.casefold() == result.away.casefold():
            errors.append("Home and Away are the same Team.")

        if date is not None and time is not None:
            result.kick_off = local_kick_off(date, time)
            if not season.starts_on <= date <= season.ends_on:
                errors.append(
                    f"{_day(date)} is outside the {season.name} Season "
                    f"({_day(season.starts_on)} to {_day(season.ends_on)})."
                )
            if result.kick_off <= now:
                errors.append("This kick-off has passed.")

        if result.pitch:
            try:
                result.pitch_obj = _find(*pitches, None, result.pitch)
            except LookupError:
                errors.append(f'Pitch "{result.pitch}" matches several pitches. Write it exactly.')
            else:
                if result.pitch_obj is None:
                    errors.append(f'The club has no pitch "{result.pitch}".')
                else:
                    result.pitch = result.pitch_obj.name

        if result.division and result.night is not None:
            try:
                result.division_obj = _find(*divisions, result.night, result.division)
            except LookupError:
                errors.append(
                    f'Division "{result.division}" matches several Divisions. Write it exactly.'
                )
            else:
                if result.division_obj is None:
                    errors.append(
                        f'{_weekday_name(result.night)} has no Division "{result.division}".'
                    )
                else:
                    result.division = result.division_obj.name

        if result.division_obj is not None:
            for side in ("home", "away"):
                name = getattr(result, side)
                if not name:
                    continue
                team = teams.get((result.division_obj.pk, name.strip().casefold()))
                if team is None:
                    errors.append(f'{result.division} has no Team "{name}" in {season.name}.')
                elif team.status != Team.Status.ACTIVE:
                    errors.append(f"{team.name} is still waiting for the Admin to confirm it.")
                else:
                    setattr(result, f"{side}_obj", team)
                    setattr(result, side, team.name)

        if errors:
            continue

        key = (result.division_obj.pk, result.home_obj.pk, result.away_obj.pk, result.kick_off)
        if key in seen_fixture:
            errors.append(f"This Fixture is already on row {seen_fixture[key]}.")
            continue
        seen_fixture[key] = row.number

        match = by_key.get(key)
        own = match.pk if match is not None else None
        pitch_key = (result.pitch_obj.pk, result.kick_off)
        clash = pitch_busy.get(pitch_key)
        if clash is not None and clash.pk != own:
            errors.append(
                f"{result.pitch} already has {clash.home_team.name} vs "
                f"{clash.away_team.name} at this kick-off."
            )
        elif pitch_key in seen_pitch:
            errors.append(
                f"{result.pitch} already has a Fixture at this time on row {seen_pitch[pitch_key]}."
            )
        seen_pitch.setdefault(pitch_key, row.number)
        for team in (result.home_obj, result.away_obj):
            team_key = (team.pk, result.kick_off)
            clash = team_busy.get(team_key)
            if clash is not None and clash.pk != own:
                errors.append(f"{team.name} already plays at this kick-off.")
            elif team_key in seen_team:
                errors.append(
                    f"{team.name} already plays at this time on row {seen_team[team_key]}."
                )
            seen_team.setdefault(team_key, row.number)
        if errors:
            continue

        result.existing = match
        if match is None:
            result.action = Action.CREATE
            result.friendly_eligible = bool(flag)
            continue
        if match.pitch_id != result.pitch_obj.pk:
            errors.append(
                f"This Fixture is on {match.pitch.name}. The import doesn't move Fixtures: "
                "reschedule it instead."
            )
            continue
        result.friendly_eligible = match.friendly_eligible if flag is None else flag
        if result.friendly_eligible != match.friendly_eligible:
            if match.status not in lifecycle.FLAG_OPEN:
                errors.append(
                    f"This Fixture is {Fixture.Status(match.status).label.lower()}, so "
                    "Friendly eligible can't change."
                )
                continue
            result.changes.append("friendly_eligible")
        result.action = Action.UPDATE if result.changes else Action.UNCHANGED

    for result in planned:
        if result.errors:
            result.action = Action.ERROR
    return Plan(planned)


def apply(club: Club, season: Season, import_plan: Plan, admin) -> None:
    """Writes a plan with no errors. Call it in a transaction."""
    if import_plan.has_errors:
        raise ValueError("A plan with errors can't be applied.")
    now = timezone.now()
    for row in import_plan.rows:
        if row.action == Action.CREATE:
            fixture = Fixture.objects.create(
                club=club,
                season=season,
                division=row.division_obj,
                pitch=row.pitch_obj,
                kick_off=row.kick_off,
                home_team=row.home_obj,
                away_team=row.away_obj,
                friendly_eligible=row.friendly_eligible,
                friendly_eligible_changed_by=admin if row.friendly_eligible else None,
                friendly_eligible_changed_at=now if row.friendly_eligible else None,
                status=Fixture.Status.SCHEDULED,
            )
            lifecycle.log_creation(fixture, cause=FixtureTransition.Cause.IMPORT, actor=admin)
        elif row.action == Action.UPDATE:
            lifecycle.set_friendly_eligible(row.existing, row.friendly_eligible, admin)


def template_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fixtures"
    sheet.append([headers[0] for headers in COLUMNS.values()])
    for row in TEMPLATE_ROWS:
        sheet.append(list(row))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for letter, width in zip("ABCDEFGH", (12, 8, 12, 14, 12, 18, 18, 16), strict=True):
        sheet.column_dimensions[letter].width = width

    help_sheet = workbook.create_sheet("How to fill it in")
    for row in TEMPLATE_HELP:
        help_sheet.append(list(row))
    help_sheet["A1"].font = help_sheet["B1"].font = Font(bold=True)
    help_sheet.column_dimensions["A"].width = 18
    help_sheet.column_dimensions["B"].width = 90

    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def template_csv() -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow([headers[0] for headers in COLUMNS.values()])
    for row in TEMPLATE_ROWS:
        writer.writerow(row)
    return out.getvalue().encode("utf-8-sig")
