"""Importing a Season's Fixtures from a spreadsheet (FR-53)."""

import datetime as dt
import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from accounts.models import ClubRole
from clubs.models import Field, Season
from fixtures import imports
from fixtures.models import Fixture
from teams.models import Team

from .conftest import next_weekday

HEADER = ["Date", "Time", "Night", "Division", "Pitch", "Home", "Away", "Friendly eligible"]


def url(club, step):
    return reverse(f"api-v1:club-import-fixtures-{step}", kwargs={"club_id": club.pk})


def csv_file(*rows, header=HEADER):
    lines = [",".join(header), *(",".join(f'"{cell}"' for cell in row) for row in rows)]
    return SimpleUploadedFile("fixtures.csv", "\n".join(lines).encode(), content_type="text/csv")


def xlsx_file(*rows, header=HEADER):
    workbook = Workbook()
    workbook.active.append(header)
    for row in rows:
        workbook.active.append(list(row))
    out = io.BytesIO()
    workbook.save(out)
    return SimpleUploadedFile("fixtures.xlsx", out.getvalue())


def upload(client, club, step, season, file):
    return client.post(url(club, step), {"season": season.pk, "file": file}, format="multipart")


@pytest.fixture
def admin(admin_of, club_a):
    return admin_of(club_a)


@pytest.fixture
def setup(lions_and_boys, make_team, make_user, pitch, next_tuesday):
    """Red Lions, Old Boys and City Rovers in Division 2; Pitch 3; next Tuesday's date."""
    rovers = make_team("City Rovers", make_user("cap@example.com", position="outfield"))
    return {"teams": (*lions_and_boys, rovers), "pitch": pitch, "date": next_tuesday.isoformat()}


def row(date, time="19:40", night="Tuesday", division="Division 2", pitch="Pitch 3",
        home="Red Lions", away="Old Boys", friendly=""):  # fmt: skip
    return (date, time, night, division, pitch, home, away, friendly)


def errors_of(answer):
    return [r["errors"] for r in answer.json()["rows"]]


# Preview and commit


@pytest.mark.django_db
def test_preview_says_what_would_happen_and_writes_nothing(admin, club_a, season, setup):
    date = setup["date"]
    file = csv_file(
        row(date, friendly="yes"),
        row(date, "20:30", night="", home="city rovers", away="RED LIONS"),
    )

    answer = upload(admin, club_a, "preview", season, file)

    assert answer.status_code == 200, answer.json()
    body = answer.json()
    assert body["committed"] is False
    assert body["summary"] == {
        "rows": 2,
        "fixtures_created": 2,
        "fixtures_updated": 0,
        "unchanged": 0,
        "errors": 0,
    }
    first, second = body["rows"]
    # 19:40 in South Africa is 17:40 UTC.
    kick_off = dt.datetime.fromisoformat(first["kick_off"])
    assert kick_off.astimezone(dt.UTC).isoformat() == f"{date}T17:40:00+00:00"
    assert first == first | {
        "row": 2,
        "night": 1,
        "division": "Division 2",
        "pitch": "Pitch 3",
        "home": "Red Lions",
        "away": "Old Boys",
        "friendly_eligible": True,
        "action": "create",
        "errors": [],
    }
    assert (second["night"], second["home"], second["away"]) == (1, "City Rovers", "Red Lions")
    assert second["friendly_eligible"] is False
    assert Fixture.objects.count() == 0


@pytest.mark.django_db
def test_commit_creates_scheduled_fixtures_and_logs_the_import(admin, club_a, season, setup):
    file = csv_file(row(setup["date"], friendly="yes"))

    answer = upload(admin, club_a, "commit", season, file)

    assert answer.status_code == 200
    assert answer.json()["committed"] is True
    fixture = Fixture.objects.get()
    lions, boys, _ = setup["teams"]
    assert (fixture.club, fixture.season, fixture.home_team, fixture.away_team) == (
        club_a,
        season,
        lions,
        boys,
    )
    assert fixture.status == Fixture.Status.SCHEDULED
    assert fixture.friendly_eligible is True
    assert fixture.friendly_eligible_changed_by == admin.user
    local = timezone.localtime(fixture.kick_off)
    assert (local.date().isoformat(), local.strftime("%H:%M")) == (setup["date"], "19:40")
    log = fixture.transitions.get()
    assert (log.from_status, log.to_status, log.cause, log.actor) == (
        "",
        "scheduled",
        "import",
        admin.user,
    )


@pytest.mark.django_db
def test_reimporting_the_same_file_creates_nothing_twice(admin, club_a, season, setup):
    rows = [row(setup["date"]), row(setup["date"], "20:30", home="City Rovers", away="Old Boys")]
    assert upload(admin, club_a, "commit", season, csv_file(*rows)).status_code == 200

    answer = upload(admin, club_a, "commit", season, csv_file(*rows))

    assert answer.status_code == 200
    assert answer.json()["summary"]["unchanged"] == 2
    assert Fixture.objects.count() == 2


@pytest.mark.django_db
def test_reimporting_can_change_friendly_eligible_only(admin, club_a, season, setup):
    upload(admin, club_a, "commit", season, csv_file(row(setup["date"])))

    answer = upload(admin, club_a, "commit", season, csv_file(row(setup["date"], friendly="Y")))

    assert answer.json()["rows"][0]["action"] == "update"
    assert answer.json()["rows"][0]["changes"] == ["friendly_eligible"]
    assert Fixture.objects.get().friendly_eligible is True
    # A blank cell keeps what the Fixture has.
    again = upload(admin, club_a, "commit", season, csv_file(row(setup["date"])))
    assert again.json()["rows"][0]["action"] == "unchanged"
    assert Fixture.objects.get().friendly_eligible is True


@pytest.mark.django_db
def test_the_import_never_moves_a_fixture_to_another_pitch(admin, club_a, season, setup):
    Field.objects.create(club=club_a, name="Pitch 4")
    upload(admin, club_a, "commit", season, csv_file(row(setup["date"])))

    answer = upload(admin, club_a, "preview", season, csv_file(row(setup["date"], pitch="Pitch 4")))

    assert "reschedule it instead" in errors_of(answer)[0][0]


@pytest.mark.django_db
def test_one_bad_row_and_nothing_is_saved(admin, club_a, season, setup):
    file = csv_file(row(setup["date"]), row(setup["date"], "20:30", home="Nobody FC"))

    answer = upload(admin, club_a, "commit", season, file)

    assert answer.status_code == 400
    assert answer.json()["committed"] is False
    assert answer.json()["summary"]["errors"] == 1
    assert Fixture.objects.count() == 0


@pytest.mark.django_db
def test_reads_excel_dates_and_times(admin, club_a, season, setup):
    date = dt.date.fromisoformat(setup["date"])
    file = xlsx_file(
        (date, dt.time(19, 40), "Tue", "Division 2", "Pitch 3", "Red Lions", "Old Boys", "no"),
        (
            dt.datetime(date.year, date.month, date.day),  # noqa: DTZ001 (Excel has no zone)
            dt.datetime(1899, 12, 30, 20, 30),  # noqa: DTZ001
            "",
            "Division 2",
            "pitch 3",
            "City Rovers",
            "Old Boys",
            None,
        ),
    )

    answer = upload(admin, club_a, "preview", season, file)

    assert errors_of(answer) == [[], []]


@pytest.mark.django_db
@pytest.mark.parametrize(
    "text",
    [
        "{iso}",
        "{d}/{m}/{y}",
        "{d}.{m}.{y}",
        "{d} {mon} {y}",
        "{wd} {d} {mon} {y}",
        "{d} {month} {y}",
    ],
)
def test_reads_south_african_date_forms(text):
    date = next_weekday(1)
    written = text.format(
        iso=date.isoformat(),
        d=date.day,
        m=date.month,
        y=date.year,
        mon=date.strftime("%b"),
        month=date.strftime("%B"),
        wd="Tue",
    )
    assert imports.parse_date(written) == date


@pytest.mark.parametrize("text", ["10/27/2026", "27 Oct", "2026-02-30", "Wed 27 Oct 2026", "soon"])
def test_refuses_dates_it_cant_be_sure_of(text):
    with pytest.raises(ValueError):
        imports.parse_date(text)


@pytest.mark.parametrize(
    ("text", "hour", "minute"),
    [("19:40", 19, 40), ("7:05", 7, 5), ("19h40", 19, 40), ("19:40:00", 19, 40)],
)
def test_reads_24_hour_times(text, hour, minute):
    assert imports.parse_time(text) == dt.time(hour, minute)


@pytest.mark.parametrize("text", ["7:40 pm", "24:00", "19:60", "0.8194444"])
def test_refuses_other_times(text):
    with pytest.raises(ValueError):
        imports.parse_time(text)


# Row errors


@pytest.mark.django_db
def test_every_row_problem_is_reported(admin, club_a, season, setup, make_team, make_user):
    make_team("Waiting FC", make_user("w@example.com"), status=Team.Status.REQUESTED)
    date = setup["date"]
    wednesday = (dt.date.fromisoformat(date) + dt.timedelta(days=1)).isoformat()
    past = next_weekday(1, after_days=-14).isoformat()
    after_season = next_weekday(1, after_days=70).isoformat()
    file = csv_file(
        row("", ""),
        row(date, division="Division 9"),
        row(date, pitch="Pitch 9"),
        row(date, home="Old Boys"),
        row(date, home="Waiting FC"),
        row(wednesday),
        row(past),
        row(after_season),
        row(date, friendly="maybe"),
        row(date, night="Funday"),
    )

    answer = upload(admin, club_a, "preview", season, file)

    errors = errors_of(answer)
    assert errors[0] == ["Date is empty.", "Time is empty."]
    assert errors[1] == ['Tuesday has no Division "Division 9".']
    assert errors[2] == ['The club has no pitch "Pitch 9".']
    assert errors[3] == ["Home and Away are the same Team."]
    assert errors[4] == ["Waiting FC is still waiting for the Admin to confirm it."]
    assert errors[5][0].startswith("Night is Tuesday, but")
    assert "This kick-off has passed." in errors[6]
    assert "outside the Spring Season" in errors[7][0]
    assert errors[8] == ["Friendly eligible must be yes, no or blank."]
    assert errors[9] == ['Night "Funday" isn\'t a weekday.']
    assert all(r["action"] == "error" for r in answer.json()["rows"])


@pytest.mark.django_db
def test_a_team_of_another_division_or_season_isnt_found(
    admin, club_a, season, setup, make_team, make_user, friday_division
):
    other_season = Season.objects.create(
        club=club_a, name="Winter", starts_on=season.starts_on, ends_on=season.ends_on
    )
    make_team("Ghosts", make_user("g@example.com"), season=other_season)
    make_team("Fridays", make_user("f@example.com"), division=friday_division)

    answer = upload(
        admin,
        club_a,
        "preview",
        season,
        csv_file(row(setup["date"], home="Ghosts"), row(setup["date"], away="Fridays")),
    )

    assert errors_of(answer) == [
        ['Division 2 has no Team "Ghosts" in Spring.'],
        ['Division 2 has no Team "Fridays" in Spring.'],
    ]


@pytest.mark.django_db
def test_clashes_in_the_file_are_errors(admin, club_a, season, setup):
    date = setup["date"]
    file = csv_file(
        row(date),
        row(date),
        row(date, home="City Rovers", away="Old Boys", pitch="Pitch 3"),
    )

    answer = upload(admin, club_a, "preview", season, file)

    errors = errors_of(answer)
    assert errors[0] == []
    assert errors[1] == ["This Fixture is already on row 2."]
    assert errors[2] == [
        "Pitch 3 already has a Fixture at this time on row 2.",
        "Old Boys already plays at this time on row 2.",
    ]


@pytest.mark.django_db
def test_clashes_with_scheduled_fixtures_are_errors(admin, club_a, season, setup, make_fixture):
    lions, boys, _ = setup["teams"]
    make_fixture(lions, boys)
    Field.objects.create(club=club_a, name="Pitch 4")

    answer = upload(
        admin,
        club_a,
        "preview",
        season,
        csv_file(
            row(setup["date"], home="City Rovers", away="Red Lions", pitch="Pitch 4"),
            row(setup["date"], home="City Rovers", away="Old Boys"),
        ),
    )

    assert errors_of(answer) == [
        ["Red Lions already plays at this kick-off."],
        [
            "Pitch 3 already has Red Lions vs Old Boys at this kick-off.",
            "City Rovers already plays at this time on row 2.",
            "Old Boys already plays at this kick-off.",
        ],
    ]


@pytest.mark.django_db
def test_a_cancelled_fixture_frees_its_pitch_and_teams(admin, club_a, season, setup, make_fixture):
    lions, boys, _ = setup["teams"]
    make_fixture(lions, boys, status=Fixture.Status.CANCELLED)

    answer = upload(
        admin,
        club_a,
        "commit",
        season,
        csv_file(row(setup["date"], home="City Rovers", away="Old Boys")),
    )

    assert answer.status_code == 200
    assert Fixture.objects.count() == 2


# The Season and the file


@pytest.mark.django_db
def test_the_season_must_be_this_clubs_and_not_ended(admin, club_a, club_b, setup):
    other = Season.objects.create(
        club=club_b, name="B", starts_on=dt.date(2026, 1, 1), ends_on=dt.date(2099, 1, 1)
    )
    ended = Season.objects.create(
        club=club_a, name="Old", starts_on=dt.date(2020, 1, 1), ends_on=dt.date(2020, 6, 1)
    )
    for season, message in [(other, "another club"), (ended, "has ended")]:
        answer = upload(admin, club_a, "preview", season, csv_file(row(setup["date"])))
        assert answer.status_code == 400
        assert message in answer.json()["season"][0]


@pytest.mark.django_db
def test_a_file_without_the_columns_is_refused(admin, club_a, season, setup):
    answer = upload(admin, club_a, "preview", season, csv_file(("x",), header=["Date"]))

    assert answer.status_code == 400
    assert '"Time"' in answer.json()["file"][0]


# Who may import


@pytest.mark.django_db
def test_only_this_clubs_admins_may_import(
    api, admin_of, club_a, club_b, season, setup, make_user, give_role
):
    referee = make_user("r@example.com")
    give_role(referee, club_a, ClubRole.Role.REFEREE)
    others = [admin_of(club_b)]
    for user in [referee, setup["teams"][0].captain]:
        client = type(api)()
        client.force_authenticate(user)
        others.append(client)

    for client in others:
        answer = upload(client, club_a, "commit", season, csv_file(row(setup["date"])))
        assert answer.status_code == 403
    assert upload(api, club_a, "preview", season, csv_file(row(setup["date"]))).status_code == 401
    assert api.get(url(club_a, "template")).status_code == 401
    assert Fixture.objects.count() == 0


@pytest.mark.django_db
def test_the_templates_download_and_import_cleanly(admin, club_a):
    xlsx = admin.get(url(club_a, "template"))
    csv = admin.get(url(club_a, "template"), {"file_type": "csv"})

    assert xlsx.status_code == 200
    assert "fixtures.xlsx" in xlsx["Content-Disposition"]
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert [c.value for c in workbook.worksheets[0][1]] == HEADER
    assert workbook.sheetnames[1] == "How to fill it in"
    assert csv.status_code == 200
    assert csv.content.decode("utf-8-sig").splitlines()[0] == ",".join(HEADER)
