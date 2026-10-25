"""The league import: Seasons, Nights and Divisions, and Fixtures from one workbook (FR-39,
FR-53)."""

import datetime as dt
import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from accounts.models import ClubRole
from clubs.imports import parse_fee, parse_weekday
from clubs.models import Division, Field, Night, Season
from core import tables
from fixtures import imports
from fixtures.models import Fixture
from ledger.models import LedgerEntry
from teams.models import Team

from .conftest import next_weekday

SEASONS = ["Season", "Starts on", "Ends on"]
DIVISIONS = ["Night", "Division", "Min players", "League fee"]
FIXTURES = ["Season", "Date", "Time", "Night", "Division", "Pitch", "Home", "Away", "Friendly"]


def url(club, step):
    return reverse(f"api-v1:club-import-league-{step}", kwargs={"club_id": club.pk})


def workbook(name="league.xlsx", **sheets):
    """An .xlsx with a sheet per keyword: seasons=[rows], or seasons=(header, [rows]), or a
    sheet named as written, e.g. **{"Nights & Divisions": [...]}."""
    headers = {"seasons": SEASONS, "divisions": DIVISIONS, "fixtures": FIXTURES}
    titles = {"seasons": "Seasons", "divisions": "Nights and Divisions", "fixtures": "Fixtures"}
    book = Workbook()
    book.remove(book.active)
    for key, rows in sheets.items():
        header = headers.get(key)
        if isinstance(rows, tuple):
            header, rows = rows
        sheet = book.create_sheet(titles.get(key, key))
        sheet.append(header)
        for row in rows:
            sheet.append(list(row))
    out = io.BytesIO()
    book.save(out)
    return SimpleUploadedFile(name, out.getvalue())


def upload(client, club, step, file):
    return client.post(url(club, step), {"file": file}, format="multipart")


def fixture_row(date, time="19:40", night="Tuesday", division="Division 2", pitch="Pitch 3",
                home="Red Lions", away="Old Boys", friendly="", season="Spring"):  # fmt: skip
    return (season, date, time, night, division, pitch, home, away, friendly)


def errors_of(answer, sheet="fixtures"):
    return [r["errors"] for r in answer.json()[sheet]]


def divisions_of(club):
    return sorted(
        (d.night.weekday, d.name, d.min_players, d.league_fee)
        for d in Division.objects.filter(night__club=club).select_related("night")
    )


@pytest.fixture
def admin(admin_of, club_a):
    return admin_of(club_a)


@pytest.fixture
def setup(lions_and_boys, make_team, make_user, pitch, next_tuesday):
    """Red Lions, Old Boys and City Rovers in Division 2 (fee R 350) of Spring; Pitch 3;
    next Tuesday's date."""
    rovers = make_team("City Rovers", make_user("cap@example.com", position="outfield"))
    return {"teams": (*lions_and_boys, rovers), "pitch": pitch, "date": next_tuesday.isoformat()}


# The whole league from one file


@pytest.mark.django_db
def test_one_file_sets_up_a_new_season_with_its_divisions_teams_and_fixtures(admin, club_a, pitch):
    settings = club_a.settings
    settings.league_fee = Decimal("350")
    settings.save()
    starts = next_weekday(0)
    ends = starts + dt.timedelta(weeks=10)
    tuesday = (starts + dt.timedelta(days=1)).isoformat()
    file = lambda: workbook(  # noqa: E731
        seasons=[("Winter", starts.isoformat(), ends.isoformat())],
        divisions=[("Tuesday", "Division 2", "6", "")],
        fixtures=[
            fixture_row(tuesday, season="Winter", friendly="yes"),
            fixture_row(tuesday, "20:30", season="winter", home="City Rovers", away="red lions"),
        ],
    )

    preview = upload(admin, club_a, "preview", file())

    assert preview.status_code == 200, preview.json()
    body = preview.json()
    assert body["committed"] is False
    assert body["summary"] == {
        "rows": 4,
        "seasons_created": 1,
        "seasons_updated": 0,
        "nights_created": 1,
        "divisions_created": 1,
        "divisions_updated": 0,
        "teams_created": 3,
        "fixtures_created": 2,
        "fixtures_updated": 0,
        "unchanged": 0,
        "errors": 0,
    }
    assert body["seasons"] == [
        {
            "row": 2,
            "name": "Winter",
            "starts_on": starts.isoformat(),
            "ends_on": ends.isoformat(),
            "action": "create",
            "changes": [],
            "errors": [],
        }
    ]
    first, second = body["fixtures"]
    assert (first["season"], first["new_teams"]) == ("Winter", ["Red Lions", "Old Boys"])
    assert (second["season"], second["new_teams"]) == ("Winter", ["City Rovers"])
    assert second["away"] == "Red Lions"
    # The preview saved nothing.
    assert not Season.objects.exists()
    assert not Night.objects.exists()
    assert not Team.objects.exists()
    assert not Fixture.objects.exists()
    assert not LedgerEntry.objects.exists()

    commit = upload(admin, club_a, "commit", file())

    assert commit.status_code == 200, commit.json()
    assert commit.json()["committed"] is True
    winter = Season.objects.get(club=club_a)
    assert (winter.name, winter.starts_on, winter.ends_on) == ("Winter", starts, ends)
    assert divisions_of(club_a) == [(1, "Division 2", 6, None)]
    teams = Team.objects.filter(season=winter).order_by("name")
    assert [(t.name, t.status, t.captain, t.decided_by) for t in teams] == [
        ("City Rovers", Team.Status.ACTIVE, None, admin.user),
        ("Old Boys", Team.Status.ACTIVE, None, admin.user),
        ("Red Lions", Team.Status.ACTIVE, None, admin.user),
    ]
    assert all(not team.memberships.exists() for team in teams)
    fees = LedgerEntry.objects.filter(team__season=winter)
    assert sorted((e.team.name, e.amount, e.category) for e in fees) == [
        ("City Rovers", Decimal("350.00"), "league"),
        ("Old Boys", Decimal("350.00"), "league"),
        ("Red Lions", Decimal("350.00"), "league"),
    ]
    assert Fixture.objects.filter(season=winter).count() == 2

    again = upload(admin, club_a, "commit", file())

    assert again.status_code == 200
    assert again.json()["summary"]["unchanged"] == 4
    assert again.json()["summary"]["teams_created"] == 0
    assert Team.objects.count() == 3
    assert Fixture.objects.count() == 2
    assert LedgerEntry.objects.count() == 3


@pytest.mark.django_db
def test_any_sheet_may_be_left_out_and_other_sheets_are_ignored(admin, club_a):
    file = workbook(
        divisions=[("Tuesday", "Division 1", "", "")],
        **{"Notes": (["Anything"], [("at all",)])},
    )

    answer = upload(admin, club_a, "commit", file)

    assert answer.status_code == 200, answer.json()
    assert (answer.json()["seasons"], answer.json()["fixtures"]) == ([], [])
    assert divisions_of(club_a) == [(1, "Division 1", 5, None)]


@pytest.mark.django_db
def test_sheets_and_headers_are_matched_loosely(admin, club_a):
    file = workbook(
        **{
            " nights & divisions ": (
                ["  NIGHT ", "Division name", "Minimum players", "League fee per Team", "Notes"],
                [("Tuesday", "Division 1", 6, 350), ("Friday", "Division 2", None, 400.5)],
            ),
            "SEASON": (
                ["Name", "Start date", "End date"],
                [("Spring", "2099-01-01", "2099-03-01")],
            ),
        }
    )

    body = upload(admin, club_a, "preview", file).json()

    read = [
        (r["night"], r["division"], r["min_players"], r["league_fee"]) for r in body["divisions"]
    ]
    assert read == [(1, "Division 1", 6, "350.00"), (4, "Division 2", 5, "400.50")]
    assert body["seasons"][0]["name"] == "Spring"


@pytest.mark.django_db
def test_one_bad_row_in_any_sheet_and_nothing_is_saved(admin, club_a, setup):
    file = workbook(
        seasons=[("Winter", "2099-01-01", "2099-03-01")],
        divisions=[("Tuesday", "Division 9", "", "")],
        fixtures=[fixture_row(setup["date"], home="Nobody FC", pitch="Pitch 9")],
    )

    answer = upload(admin, club_a, "commit", file)

    assert answer.status_code == 400
    assert answer.json()["committed"] is False
    assert answer.json()["summary"]["errors"] == 1
    assert errors_of(answer) == [['The club has no pitch "Pitch 9".']]
    assert not Season.objects.filter(name="Winter").exists()
    assert not Division.objects.filter(name="Division 9").exists()
    assert not Team.objects.filter(name="Nobody FC").exists()


@pytest.mark.django_db
def test_commit_is_all_or_nothing_when_a_write_fails(admin, club_a, setup, monkeypatch):
    def fail(*args, **kwargs):
        raise RuntimeError("The database went away")

    monkeypatch.setattr(imports.lifecycle, "log_creation", fail)
    file = workbook(
        seasons=[("Winter", "2099-01-01", "2099-03-01")],
        fixtures=[fixture_row(setup["date"], home="New Stars")],
    )

    with pytest.raises(RuntimeError):
        upload(admin, club_a, "commit", file)

    assert not Season.objects.filter(name="Winter").exists()
    assert not Team.objects.filter(name="New Stars").exists()
    assert not Fixture.objects.exists()


# The file


@pytest.mark.django_db
@pytest.mark.parametrize(
    ("file", "message"),
    [
        (
            workbook(divisions=(["Night", "Min players"], [])),
            'The Nights and Divisions sheet has no "Division" column.',
        ),
        (workbook(**{"Sheet1": (["Night"], [])}), "The workbook has none of the sheets"),
        (SimpleUploadedFile("league.csv", b"Night,Division"), "Upload an Excel workbook (.xlsx)."),
        (SimpleUploadedFile("league.xlsx", b"not a workbook"), "can't be read as an Excel"),
    ],
)
def test_a_file_that_cant_be_used_is_refused_whole(admin, club_a, file, message):
    answer = upload(admin, club_a, "preview", file)

    assert answer.status_code == 400
    assert message in answer.json()["file"][0]


@pytest.mark.django_db
def test_files_are_limited_in_size_and_rows_per_sheet(admin, club_a, monkeypatch):
    monkeypatch.setattr(tables, "MAX_ROWS", 2)

    too_long = workbook(divisions=[("Tuesday", f"Division {n}", "", "") for n in range(3)])
    message = upload(admin, club_a, "preview", too_long).json()["file"][0]
    assert message == "The Nights and Divisions sheet has more than 2 rows. Import it in parts."

    monkeypatch.setattr(tables, "MAX_BYTES", 200)
    too_big = workbook(divisions=[("Tuesday", "Division 1", "", "")])
    assert upload(admin, club_a, "preview", too_big).json()["file"] == ["The file is over 5 MB."]


@pytest.mark.django_db
def test_the_template_downloads_and_imports_cleanly(admin, club_a, admin_of, club_b):
    settings = club_a.settings
    settings.league_fee = Decimal("350")
    settings.save()
    Field.objects.create(club=club_a, name="Pitch 3")

    answer = admin.get(url(club_a, "template"))

    assert answer.status_code == 200
    assert answer["Content-Disposition"] == 'attachment; filename="league-import.xlsx"'
    book = load_workbook(io.BytesIO(answer.content))
    assert book.sheetnames == ["Seasons", "Nights and Divisions", "Fixtures", "How to fill it in"]
    assert [c.value for c in book["Seasons"][1]] == SEASONS
    assert [c.value for c in book["Nights and Divisions"][1]] == DIVISIONS
    assert [c.value for c in book["Fixtures"][1]] == [*FIXTURES[:-1], "Friendly eligible"]
    template = SimpleUploadedFile("template.xlsx", answer.content)
    preview = upload(admin, club_a, "preview", template).json()
    assert preview["summary"]["errors"] == 0, preview
    assert admin_of(club_b).get(url(club_a, "template")).status_code == 403


# Seasons


@pytest.mark.django_db
def test_an_existing_season_takes_new_dates_unless_its_fixtures_fall_outside(
    admin, club_a, season, setup, make_fixture, next_tuesday
):
    make_fixture(*setup["teams"][:2])
    later = (season.ends_on + dt.timedelta(days=14)).isoformat()
    file = workbook(seasons=[("SPRING", season.starts_on.isoformat(), later)])

    grows = upload(admin, club_a, "commit", file)

    assert grows.status_code == 200, grows.json()
    assert grows.json()["seasons"][0] | {"errors": []} == grows.json()["seasons"][0]
    assert grows.json()["seasons"][0]["action"] == "update"
    assert grows.json()["seasons"][0]["changes"] == ["ends_on"]
    assert grows.json()["seasons"][0]["name"] == "Spring"
    season.refresh_from_db()
    assert season.ends_on.isoformat() == later

    too_short = (next_tuesday - dt.timedelta(days=1)).isoformat()
    shrinks = upload(
        admin,
        club_a,
        "preview",
        workbook(seasons=[("Spring", season.starts_on.isoformat(), too_short)]),
    )
    (message,) = errors_of(shrinks, "seasons")[0]
    assert message.startswith("Spring has Fixtures from ")
    assert "so it must start by" in message


@pytest.mark.django_db
def test_every_season_row_problem_is_reported(admin, club_a):
    file = workbook(
        seasons=[
            ("", "", ""),
            ("Winter", "2099-03-01", "2099-01-01"),
            ("Autumn", "soon", "2099-02-30"),
            ("S" * 61, "2099-01-01", "2099-03-01"),
            ("Summer", "2099-01-01", "2099-03-01"),
            ("summer ", "2099-01-01", "2099-03-01"),
        ]
    )

    answer = upload(admin, club_a, "preview", file)

    assert [r["row"] for r in answer.json()["seasons"]] == [3, 4, 5, 6, 7]  # Blank row skipped.
    errors = errors_of(answer, "seasons")
    assert errors[0] == ["Ends on is before Starts on."]
    assert errors[1] == [
        "Starts on must be a date, like 2026-10-27, 27/10/2026 or 27 Oct 2026.",
        "Ends on 2099-02-30 doesn't exist.",
    ]
    assert errors[2] == ["Season is longer than 60 characters."]
    assert errors[3] == []
    assert errors[4] == ["This Season is already on row 6."]


# Nights and Divisions


@pytest.mark.parametrize(
    ("text", "weekday"),
    [("Monday", 0), ("tue", 1), ("TUES", 1), ("Wednesday Night", 2), ("thurs", 3), ("Sun", 6)],
)
def test_weekdays_are_read_in_any_common_form(text, weekday):
    assert parse_weekday(text) == weekday


@pytest.mark.parametrize(
    ("text", "fee"),
    [
        ("350", "350.00"),
        ("R 350", "350.00"),
        ("R\u00a0350", "350.00"),
        ("R1 200", "1200.00"),
        ("1,200.5", "1200.50"),
    ],
)
def test_fees_are_read_as_rand(text, fee):
    assert parse_fee(text) == Decimal(fee)


@pytest.mark.django_db
def test_each_bad_division_cell_is_reported_on_its_row(admin, club_a):
    file = workbook(
        divisions=[
            ("Someday", "Division 1", "", ""),
            ("Tuesday", "", "0", "350,50"),
            ("Tu", "D" * 61, "6.5", "-10"),
        ]
    )

    errors = errors_of(upload(admin, club_a, "preview", file), "divisions")

    assert errors[0] == ['Night "Someday" isn\'t a weekday.']
    assert errors[1] == [
        "Division is empty.",
        "Min players must be a whole number of 1 or more.",
        "League fee must be an amount in rand, like 350 or 1,200.50.",
    ]
    assert errors[2] == [
        'Night "Tu" isn\'t a weekday.',
        "Division is longer than 60 characters.",
        "Min players must be a whole number of 1 or more.",
        "League fee must be an amount in rand, like 350 or 1,200.50.",
    ]


@pytest.mark.django_db
def test_re_importing_divisions_updates_values_keeps_blanks_and_leaves_others(
    admin, club_a, tuesday
):
    Division.objects.create(night=tuesday, name="Division 1", min_players=6, league_fee=350)
    Division.objects.create(night=tuesday, name="Division 2", min_players=7)

    answer = upload(
        admin, club_a, "commit", workbook(divisions=[("TUESDAY", "DIVISION 1", "", "400")])
    )

    assert answer.json()["divisions"][0]["changes"] == ["league_fee"]
    assert divisions_of(club_a) == [
        (1, "Division 1", 6, Decimal("400.00")),  # Name as it was; blank min kept.
        (1, "Division 2", 7, None),  # Not in the file: left alone.
    ]


@pytest.mark.django_db
def test_divisions_differing_only_in_case_are_matched_exactly_or_refused(admin, club_a, tuesday):
    Division.objects.create(night=tuesday, name="Division A", min_players=5)
    Division.objects.create(night=tuesday, name="division a", min_players=8)

    answer = upload(
        admin,
        club_a,
        "commit",
        workbook(divisions=[("Tuesday", "division a", "9", ""), ("Tuesday", "Division A", "", "")]),
    )
    assert answer.status_code == 200, answer.json()
    assert divisions_of(club_a) == [(1, "Division A", 5, None), (1, "division a", 9, None)]

    preview = upload(
        admin, club_a, "preview", workbook(divisions=[("Tuesday", "DIVISION A", "7", "")])
    )
    assert errors_of(preview, "divisions") == [
        ['Division matches "Division A", "division a". Write the name exactly as the club has it.']
    ]


# Fixtures


@pytest.mark.django_db
def test_fixtures_are_created_scheduled_and_the_import_logged(admin, club_a, season, setup):
    answer = upload(
        admin, club_a, "commit", workbook(fixtures=[fixture_row(setup["date"], friendly="yes")])
    )

    assert answer.status_code == 200, answer.json()
    row = answer.json()["fixtures"][0]
    # 19:40 in South Africa is 17:40 UTC.
    kick_off = dt.datetime.fromisoformat(row["kick_off"])
    assert kick_off.astimezone(dt.UTC).isoformat() == f"{setup['date']}T17:40:00+00:00"
    assert row["new_teams"] == []
    fixture = Fixture.objects.get()
    lions, boys, _ = setup["teams"]
    assert (fixture.club, fixture.season, fixture.home_team, fixture.away_team) == (
        club_a,
        season,
        lions,
        boys,
    )
    assert fixture.status == Fixture.Status.SCHEDULED
    assert fixture.friendly_eligible is True
    assert fixture.friendly_eligible_changed_by == admin.user
    local = timezone.localtime(fixture.kick_off)
    assert (local.date().isoformat(), local.strftime("%H:%M")) == (setup["date"], "19:40")
    log = fixture.transitions.get()
    assert (log.from_status, log.to_status, log.cause, log.actor) == (
        "",
        "scheduled",
        "import",
        admin.user,
    )


@pytest.mark.django_db
def test_reimporting_can_change_friendly_eligible_only(admin, club_a, season, setup):
    upload(admin, club_a, "commit", workbook(fixtures=[fixture_row(setup["date"])]))

    answer = upload(
        admin, club_a, "commit", workbook(fixtures=[fixture_row(setup["date"], friendly="Y")])
    )

    assert answer.json()["fixtures"][0]["action"] == "update"
    assert answer.json()["fixtures"][0]["changes"] == ["friendly_eligible"]
    assert Fixture.objects.get().friendly_eligible is True
    # A blank cell keeps what the Fixture has.
    again = upload(admin, club_a, "commit", workbook(fixtures=[fixture_row(setup["date"])]))
    assert again.json()["fixtures"][0]["action"] == "unchanged"
    assert Fixture.objects.get().friendly_eligible is True


@pytest.mark.django_db
def test_the_import_never_moves_a_fixture_to_another_pitch(admin, club_a, season, setup):
    Field.objects.create(club=club_a, name="Pitch 4")
    upload(admin, club_a, "commit", workbook(fixtures=[fixture_row(setup["date"])]))

    answer = upload(
        admin, club_a, "preview", workbook(fixtures=[fixture_row(setup["date"], pitch="Pitch 4")])
    )

    assert "reschedule it instead" in errors_of(answer)[0][0]


@pytest.mark.django_db
def test_reads_excel_dates_and_times(admin, club_a, season, setup):
    date = dt.date.fromisoformat(setup["date"])
    file = workbook(
        fixtures=[
            ("Spring", date, dt.time(19, 40), "Tue", "Division 2", "Pitch 3", "Red Lions",
             "Old Boys", "no"),
            (
                "Spring",
                dt.datetime(date.year, date.month, date.day),  # noqa: DTZ001 (Excel has no zone)
                dt.datetime(1899, 12, 30, 20, 30),  # noqa: DTZ001
                "",
                "Division 2",
                "pitch 3",
                "City Rovers",
                "Old Boys",
                None,
            ),
        ]
    )  # fmt: skip

    assert errors_of(upload(admin, club_a, "preview", file)) == [[], []]


@pytest.mark.parametrize(
    "text",
    [
        "{iso}",
        "{d}/{m}/{y}",
        "{d}.{m}.{y}",
        "{d} {mon} {y}",
        "{wd} {d} {mon} {y}",
        "{d} {month} {y}",
    ],
)
def test_reads_south_african_date_forms(text):
    date = next_weekday(1)
    written = text.format(
        iso=date.isoformat(),
        d=date.day,
        m=date.month,
        y=date.year,
        mon=date.strftime("%b"),
        month=date.strftime("%B"),
        wd="Tue",
    )
    assert imports.parse_date(written) == date


@pytest.mark.parametrize("text", ["10/27/2026", "27 Oct", "2026-02-30", "Wed 27 Oct 2026", "soon"])
def test_refuses_dates_it_cant_be_sure_of(text):
    with pytest.raises(ValueError):
        imports.parse_date(text)


@pytest.mark.parametrize(
    ("text", "hour", "minute"),
    [("19:40", 19, 40), ("7:05", 7, 5), ("19h40", 19, 40), ("19:40:00", 19, 40)],
)
def test_reads_24_hour_times(text, hour, minute):
    assert imports.parse_time(text) == dt.time(hour, minute)


@pytest.mark.parametrize("text", ["7:40 pm", "24:00", "19:60", "0.8194444"])
def test_refuses_other_times(text):
    with pytest.raises(ValueError):
        imports.parse_time(text)


@pytest.mark.django_db
def test_every_fixture_row_problem_is_reported(admin, club_a, season, setup, make_team, make_user):
    make_team("Waiting FC", make_user("w@example.com"), status=Team.Status.REQUESTED)
    date = setup["date"]
    wednesday = (dt.date.fromisoformat(date) + dt.timedelta(days=1)).isoformat()
    past = next_weekday(1, after_days=-14).isoformat()
    after_season = next_weekday(1, after_days=70).isoformat()
    file = workbook(
        fixtures=[
            fixture_row("", "", season=""),
            fixture_row(date, division="Division 9"),
            fixture_row(date, pitch="Pitch 9"),
            fixture_row(date, home="Old Boys"),
            fixture_row(date, home="Waiting FC"),
            fixture_row(wednesday),
            fixture_row(past),
            fixture_row(after_season),
            fixture_row(date, friendly="maybe"),
            fixture_row(date, night="Funday"),
            fixture_row(date, season="Winter"),
        ]
    )

    errors = errors_of(upload(admin, club_a, "preview", file))

    assert errors[0] == ["Season is empty.", "Date is empty.", "Time is empty."]
    assert errors[1] == ['Tuesday has no Division "Division 9".']
    assert errors[2] == ['The club has no pitch "Pitch 9".']
    assert errors[3] == ["Home and Away are the same Team."]
    assert errors[4] == ["Waiting FC is still waiting for the Admin to confirm it."]
    assert errors[5][0].startswith("Night is Tuesday, but")
    assert "This kick-off has passed." in errors[6]
    assert "outside the Spring Season" in errors[7][0]
    assert errors[8] == ["Friendly eligible must be yes, no or blank."]
    assert errors[9] == ['Night "Funday" isn\'t a weekday.']
    assert errors[10] == ['The club has no Season "Winter". Add it on the Seasons sheet.']


@pytest.mark.django_db
def test_clashes_in_the_file_and_with_scheduled_fixtures_are_errors(
    admin, club_a, season, setup, make_fixture
):
    lions, boys, _ = setup["teams"]
    make_fixture(lions, boys, time="20:30")
    date = setup["date"]
    file = workbook(
        fixtures=[
            fixture_row(date),
            fixture_row(date),
            fixture_row(date, home="City Rovers", away="Old Boys"),
            fixture_row(date, "20:30", home="City Rovers", away="Red Lions"),
            fixture_row(date, "21:20", home="New Stars", away="City Rovers"),
            fixture_row(date, "21:20", home="new stars", away="Old Boys", pitch="Pitch 3"),
        ]
    )

    errors = errors_of(upload(admin, club_a, "preview", file))

    assert errors[0] == []
    assert errors[1] == ["This Fixture is already on row 2."]
    assert errors[2] == [
        "Pitch 3 already has a Fixture at this time on row 2.",
        "Old Boys already plays at this time on row 2.",
    ]
    assert errors[3] == [
        "Pitch 3 already has Red Lions vs Old Boys at this kick-off.",
        "Red Lions already plays at this kick-off.",
    ]
    assert errors[4] == []
    # A Team the import adds is the same Team on every row, so it clashes too.
    assert errors[5] == [
        "Pitch 3 already has a Fixture at this time on row 6.",
        "New Stars already plays at this time on row 6.",
    ]


@pytest.mark.django_db
def test_a_cancelled_fixture_frees_its_pitch_and_teams(admin, club_a, season, setup, make_fixture):
    lions, boys, _ = setup["teams"]
    make_fixture(lions, boys, status=Fixture.Status.CANCELLED)

    answer = upload(
        admin,
        club_a,
        "commit",
        workbook(fixtures=[fixture_row(setup["date"], home="City Rovers", away="Old Boys")]),
    )

    assert answer.status_code == 200, answer.json()
    assert Fixture.objects.count() == 2


# Teams the import adds


@pytest.mark.django_db
def test_a_team_of_another_division_or_season_is_a_new_team(
    admin, club_a, season, setup, make_team, make_user, friday_division
):
    other_season = Season.objects.create(
        club=club_a, name="Winter", starts_on=season.starts_on, ends_on=season.ends_on
    )
    ghosts = make_team("Ghosts", make_user("g@example.com"), season=other_season)
    fridays = make_team("Fridays", make_user("f@example.com"), division=friday_division)

    answer = upload(
        admin,
        club_a,
        "commit",
        workbook(fixtures=[fixture_row(setup["date"], home="Ghosts", away="Fridays")]),
    )

    assert answer.status_code == 200, answer.json()
    assert answer.json()["fixtures"][0]["new_teams"] == ["Ghosts", "Fridays"]
    added = Team.objects.filter(season=season, division=setup["teams"][0].division)
    assert {"Ghosts", "Fridays"} <= set(added.values_list("name", flat=True))
    assert Team.objects.filter(name="Ghosts").exclude(pk=ghosts.pk).get().captain is None
    assert Team.objects.filter(name="Fridays").exclude(pk=fridays.pk).count() == 1


@pytest.mark.django_db
def test_no_team_is_added_without_a_league_fee_or_past_the_divisions_max(
    admin, club_a, season, setup, friday_division, pitch
):
    division = setup["teams"][0].division
    division.max_teams = 4
    division.save()
    friday = next_weekday(4).isoformat()

    answer = upload(
        admin,
        club_a,
        "preview",
        workbook(
            fixtures=[
                fixture_row(setup["date"], home="New Stars", away="Red Lions"),
                fixture_row(setup["date"], "20:30", home="Hat Trick FC", away="Old Boys"),
                fixture_row(friday, night="Friday", division="Division 1", home="A", away="B"),
            ]
        ),
    )

    errors = errors_of(answer)
    assert errors[0] == []
    assert errors[1] == [
        "Division 2 has all the Teams it takes in Spring, so the import can't add Hat Trick FC."
    ]
    club_a.settings.league_fee = None
    club_a.settings.save()
    no_fee = upload(
        admin,
        club_a,
        "preview",
        workbook(
            fixtures=[
                fixture_row(friday, night="Friday", division="Division 1", home="A", away="B")
            ]
        ),
    )
    assert errors_of(no_fee)[0] == [
        "The club hasn't set the league fee for Division 1, so the import can't add A.",
        "The club hasn't set the league fee for Division 1, so the import can't add B.",
    ]


# Who may import


@pytest.mark.django_db
@pytest.mark.parametrize("step", ["preview", "commit"])
def test_only_this_clubs_admins_may_import(
    api, admin_of, club_a, club_b, season, setup, make_user, give_role, step
):
    referee = make_user("r@example.com")
    give_role(referee, club_a, ClubRole.Role.REFEREE)
    others = [admin_of(club_b)]
    for user in [referee, setup["teams"][0].captain]:
        client = type(api)()
        client.force_authenticate(user)
        others.append(client)
    file = lambda: workbook(seasons=[("Winter", "2099-01-01", "2099-03-01")])  # noqa: E731

    for client in others:
        assert upload(client, club_a, step, file()).status_code == 403
    assert upload(api, club_a, step, file()).status_code == 401
    assert api.get(url(club_a, "template")).status_code == 401
    assert not Season.objects.filter(name="Winter").exists()


@pytest.mark.django_db
def test_an_import_only_reads_and_touches_its_own_club(admin, club_a, club_b, season):
    Season.objects.create(
        club=club_b, name="Winter", starts_on=season.starts_on, ends_on=season.ends_on
    )
    other = Night.objects.create(club=club_b, weekday=Night.Weekday.TUESDAY)
    Division.objects.create(night=other, name="Division 1", min_players=9)

    answer = upload(
        admin,
        club_a,
        "commit",
        workbook(
            seasons=[("Winter", "2099-01-01", "2099-03-01")],
            divisions=[("Tuesday", "Division 1", "", "")],
        ),
    )

    assert answer.json()["seasons"][0]["action"] == "create"
    assert answer.json()["divisions"][0]["action"] == "create"
    assert divisions_of(club_b) == [(1, "Division 1", 9, None)]
    assert Season.objects.get(club=club_b, name="Winter").starts_on == season.starts_on
