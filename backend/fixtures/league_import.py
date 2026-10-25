"""The league import (FR-39, FR-53): one Excel workbook with a club's Seasons, its Nights
and Divisions, and its Fixtures, a sheet each. Any sheet may be left out.

The sheets are checked in that order, each against what the sheets before it add, so a new
Season or Division can have its Fixtures in the same file. Seasons are matched on their
name, ignoring case and outer spaces (docs/DECISIONS.md):

- a new Season is created;
- an existing Season takes the file's dates, unless one of its Fixtures would fall outside
  them;
- Seasons that aren't in the file are left alone.

The Nights and Divisions sheet is clubs/imports.py, the Fixtures sheet fixtures/imports.py.

run() writes each sheet's rows that have no errors before checking the next sheet, so it
must run in a transaction that is rolled back unless the outcome has no errors and the
Admin is committing. A preview always rolls back.
"""

import dataclasses
import datetime as dt
import io
from dataclasses import dataclass, field

from django.db.models import Max, Min
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font

from clubs import imports as division_imports
from clubs.models import Club, Season
from core.tables import ImportAction as Action
from core.tables import Row, Sheet

from . import imports as fixture_imports
from .imports import _by_name, _day, _find, parse_date
from .models import Fixture

SEASON_COLUMNS = {
    "name": ("Season", "Name", "Season name"),
    "starts_on": ("Starts on", "Starts", "Start", "Start date"),
    "ends_on": ("Ends on", "Ends", "End", "End date"),
}
SEASON_HELP = [
    ("Season", "The Season's name, e.g. Summer 2027. Required."),
    ("Starts on", "Its first day: 2027-01-12, 12/01/2027 or 12 Jan 2027. Required."),
    ("Ends on", "Its last day, written the same way. Required."),
    ("", "A Season the club has takes the file's dates, unless its Fixtures fall outside."),
]

SHEETS = {
    "seasons": Sheet(("Seasons", "Season"), SEASON_COLUMNS, frozenset(SEASON_COLUMNS)),
    "divisions": Sheet(
        ("Nights and Divisions", "Nights & Divisions", "Divisions", "Nights"),
        division_imports.COLUMNS,
        frozenset(division_imports.REQUIRED),
    ),
    "fixtures": Sheet(
        ("Fixtures", "Fixture list"), fixture_imports.COLUMNS, frozenset(fixture_imports.REQUIRED)
    ),
}
NOTES = [
    "Leave out any sheet you don't need.",
    "Sheets are read in order: Seasons, then Nights and Divisions, then Fixtures. So a new "
    "Season or Division can have its Fixtures in the same file.",
    "Nothing is imported until every row is right.",
    "Re-importing the same file creates nothing twice.",
    "Nothing missing from the file is deleted.",
]

NAME_MAX = Season._meta.get_field("name").max_length


@dataclass
class PlannedSeason:
    row: int
    name: str = ""
    starts_on: dt.date | None = None
    ends_on: dt.date | None = None
    action: str = Action.ERROR
    changes: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    existing: Season | None = None


@dataclass
class SeasonPlan:
    rows: list[PlannedSeason]

    @property
    def has_errors(self) -> bool:
        return any(row.errors for row in self.rows)

    @property
    def summary(self) -> dict[str, int]:
        def count(action):
            return sum(1 for row in self.rows if row.action == action)

        return {
            "seasons_created": count(Action.CREATE),
            "seasons_updated": count(Action.UPDATE),
            "unchanged": count(Action.UNCHANGED),
            "errors": count(Action.ERROR),
        }


def plan_seasons(club: Club, rows: list[Row]) -> SeasonPlan:
    seasons = _by_name(club.seasons.all(), lambda s: (None, s.name.strip()))
    # The South African dates of each Season's first and last Fixture.
    spans = {
        span["season"]: (timezone.localdate(span["first"]), timezone.localdate(span["last"]))
        for span in Fixture.objects.filter(club=club)
        .values("season")
        .annotate(first=Min("kick_off"), last=Max("kick_off"))
    }
    seen: dict[tuple, int] = {}
    planned = []

    for row in rows:
        result = PlannedSeason(row=row.number)
        planned.append(result)
        cells = row.cells
        errors = result.errors

        result.name = cells.get("name", "")
        if not result.name:
            errors.append("Season is empty.")
        elif len(result.name) > NAME_MAX:
            errors.append(f"Season is longer than {NAME_MAX} characters.")
        for key, label in [("starts_on", "Starts on"), ("ends_on", "Ends on")]:
            if not cells.get(key):
                errors.append(f"{label} is empty.")
                continue
            try:
                setattr(result, key, parse_date(cells[key]))
            except ValueError as error:
                errors.append(str(error).replace("Date", label, 1))
        if result.starts_on and result.ends_on and result.ends_on < result.starts_on:
            errors.append("Ends on is before Starts on.")
        if errors:
            continue

        try:
            existing = _find(*seasons, None, result.name)
        except LookupError:
            errors.append(f'Season "{result.name}" matches several Seasons. Write it exactly.')
            continue
        key = ("season", existing.pk) if existing else ("new", result.name.casefold())
        if key in seen:
            errors.append(f"This Season is already on row {seen[key]}.")
            continue
        seen[key] = row.number

        result.existing = existing
        if existing is None:
            result.action = Action.CREATE
            continue
        result.name = existing.name
        result.changes = [
            key
            for key in ("starts_on", "ends_on")
            if getattr(result, key) != getattr(existing, key)
        ]
        span = spans.get(existing.pk)
        if result.changes and span and not result.starts_on <= span[0] <= span[1] <= result.ends_on:
            first, last = span
            errors.append(
                f"{existing.name} has Fixtures from {_day(first)} to {_day(last)}, so it must "
                f"start by {_day(first)} and end on {_day(last)} or later."
            )
            continue
        result.action = Action.UPDATE if result.changes else Action.UNCHANGED

    return SeasonPlan(planned)


def apply_seasons(club: Club, season_plan: SeasonPlan) -> None:
    """Writes a plan with no errors. Call it in a transaction."""
    if season_plan.has_errors:
        raise ValueError("A plan with errors can't be applied.")
    for row in season_plan.rows:
        if row.action == Action.CREATE:
            Season.objects.create(
                club=club, name=row.name, starts_on=row.starts_on, ends_on=row.ends_on
            )
        elif row.action == Action.UPDATE:
            season = row.existing
            season.starts_on, season.ends_on = row.starts_on, row.ends_on
            season.save(update_fields=["starts_on", "ends_on"])


@dataclass
class Outcome:
    seasons: SeasonPlan
    divisions: division_imports.Plan
    fixtures: fixture_imports.Plan

    @property
    def plans(self):
        return (self.seasons, self.divisions, self.fixtures)

    @property
    def has_errors(self) -> bool:
        return any(plan.has_errors for plan in self.plans)

    @property
    def summary(self) -> dict[str, int]:
        seasons = self.seasons.summary
        divisions = self.divisions.summary
        fixtures = self.fixtures.summary
        return {
            "rows": sum(len(plan.rows) for plan in self.plans),
            "seasons_created": seasons["seasons_created"],
            "seasons_updated": seasons["seasons_updated"],
            "nights_created": divisions["nights_created"],
            "divisions_created": divisions["divisions_created"],
            "divisions_updated": divisions["divisions_updated"],
            "teams_created": fixtures["teams_created"],
            "fixtures_created": fixtures["fixtures_created"],
            "fixtures_updated": fixtures["fixtures_updated"],
            "unchanged": sum(plan.summary["unchanged"] for plan in self.plans),
            "errors": sum(plan.summary["errors"] for plan in self.plans),
        }


def _without_errors(plan):
    return dataclasses.replace(plan, rows=[row for row in plan.rows if not row.errors])


def run(club: Club, sheets: dict[str, list[Row]], admin) -> Outcome:
    """Checks and writes each sheet in turn. Call it in a transaction; see the module."""
    seasons = plan_seasons(club, sheets.get("seasons", []))
    apply_seasons(club, _without_errors(seasons))
    divisions = division_imports.plan(club, sheets.get("divisions", []))
    division_imports.apply(club, _without_errors(divisions))
    fixtures = fixture_imports.plan(club, sheets.get("fixtures", []))
    fixture_imports.apply(club, _without_errors(fixtures), admin)
    return Outcome(seasons, divisions, fixtures)


def _example_rows(today: dt.date) -> dict[str, list[tuple]]:
    """Rows that import as they are, given a league fee and a pitch named Pitch 3: a Season
    from next Monday and Fixtures on its first Tuesday."""
    starts = today + dt.timedelta(days=7 - today.weekday())
    ends = starts + dt.timedelta(weeks=12, days=-1)
    tuesday = (starts + dt.timedelta(days=1)).isoformat()
    name = f"Season {starts:%b %Y}"
    match = (name, tuesday)
    where = ("Tuesday", "Division 2", "Pitch 3")
    return {
        "seasons": [(name, starts.isoformat(), ends.isoformat())],
        "divisions": list(division_imports.TEMPLATE_ROWS),
        "fixtures": [
            (*match, "19:40", *where, "Red Lions", "Old Boys", "yes"),
            (*match, "20:30", *where, "City Rovers", "Hat Trick FC", ""),
        ],
    }


def template_xlsx() -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    examples = _example_rows(timezone.localdate())
    for key, sheet in SHEETS.items():
        worksheet = workbook.create_sheet(sheet.name)
        worksheet.append([headers[0] for headers in sheet.columns.values()])
        for row in examples[key]:
            worksheet.append(list(row))
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            worksheet.column_dimensions[cell.column_letter].width = max(14, len(cell.value) + 4)

    help_sheet = workbook.create_sheet("How to fill it in")
    help_sheet.append(["Sheet", "Column", "What to put in it"])
    for cell in help_sheet[1]:
        cell.font = Font(bold=True)
    for key, lines in [
        ("seasons", SEASON_HELP),
        ("divisions", division_imports.TEMPLATE_HELP),
        ("fixtures", fixture_imports.TEMPLATE_HELP),
    ]:
        for column, text in lines:
            help_sheet.append([SHEETS[key].name, column, text])
    for note in NOTES:
        help_sheet.append(["", "", note])
    help_sheet.column_dimensions["A"].width = 22
    help_sheet.column_dimensions["B"].width = 18
    help_sheet.column_dimensions["C"].width = 90

    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
