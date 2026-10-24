"""The accountant export (FR-51): a Club's ledger entries for any date range, as CSV or Excel.

One row per ledger entry recorded in the range (South African days, both inclusive),
oldest first. Nothing here writes: corrections show as their own rows (FR-50), naming the
entry they correct. Columns and their meaning: docs/DECISIONS.md, chunk 28.

- **Amount** is the entry's amount as recorded: rand, always above zero.
- **Money in** and **Money out** are only card and cash the club actually took in or paid
  back: a payment is money in, a refund money out, and a correction of either goes the
  other way. Charges, credits and payments from credit move no money, so both are empty.

Text a spreadsheet would read as a formula starts with an apostrophe (CSV injection).
"""

import csv
import datetime as dt
import io
from dataclasses import dataclass
from decimal import Decimal

from django.db.models import QuerySet
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font

from clubs.models import Club
from ledger.models import LedgerEntry

Kind = LedgerEntry.Kind
Method = LedgerEntry.Method

COLUMNS = [
    "Entry",
    "Date",
    "Time",
    "Kind",
    "Category",
    "Method",
    "Amount (R)",
    "Money in (R)",
    "Money out (R)",
    "Account",
    "Team",
    "Player or Booking User",
    "Season",
    "Night",
    "Division",
    "Fixture",
    "Booking",
    "Corrects entry",
    "Credit from entry",
    "Reason",
    "Recorded by",
]
MONEY_COLUMNS = {"Amount (R)", "Money in (R)", "Money out (R)"}
WIDTHS = {"Reason": 40, "Team": 22, "Player or Booking User": 24, "Recorded by": 20}

# What the About sheet says each less obvious column means.
ABOUT = [
    ("Date, Time", "When the entry was recorded, South African time (Africa/Johannesburg)."),
    ("Amount (R)", "The entry's amount in rand, always above zero. Kind says which way it goes."),
    (
        "Money in (R), Money out (R)",
        "Card and cash the club took in (payments) or paid back (refunds). A correction of "
        "either goes the other way. Charges, credits and payments from credit move no money.",
    ),
    ("Account", "Whose account the entry is on: a Team, a Player or a Booking User."),
    ("Corrects entry", "For a correction (Reversal): the entry it undoes."),
    ("Credit from entry", "For credit given back: the payment it came from."),
    ("Recorded by", "Who recorded it; League App when the app did."),
]

FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")
CASH_MOVES = {Method.CARD, Method.CASH}


def safe_text(value: str) -> str:
    """Text as a spreadsheet shows it, never run as a formula. Control characters a
    worksheet can't hold (a pasted vertical tab in a reason) are dropped, so one odd name
    or reason can't fail the whole export."""
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    return f"'{value}" if value.startswith(FORMULA_STARTS) else value


@dataclass(frozen=True)
class Period:
    starts_on: dt.date
    ends_on: dt.date

    def window(self) -> tuple[dt.datetime, dt.datetime]:
        """The period's South African days as a [start, end) range of moments."""
        tz = timezone.get_default_timezone()
        start = dt.datetime.combine(self.starts_on, dt.time.min, tzinfo=tz)
        end = dt.datetime.combine(self.ends_on + dt.timedelta(days=1), dt.time.min, tzinfo=tz)
        return start, end

    @property
    def name(self) -> str:
        return f"ledger-{self.starts_on.isoformat()}-to-{self.ends_on.isoformat()}"


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def entries(club: Club, period: Period) -> QuerySet[LedgerEntry]:
    start, end = period.window()
    return (
        LedgerEntry.objects.filter(club=club, created_at__gte=start, created_at__lt=end)
        .select_related(
            "team__season",
            "team__division__night",
            "user",
            "fixture__season",
            "fixture__division__night",
            "fixture__home_team",
            "fixture__away_team",
            "reverses",
            "recorded_by",
        )
        .order_by("created_at", "pk")
    )


def money_in_out(entry: LedgerEntry) -> tuple[Decimal | None, Decimal | None]:
    """Card and cash taken in and paid back by this entry: (in, out)."""
    if entry.method not in CASH_MOVES:
        return None, None
    kind = entry.kind
    if kind == Kind.REVERSAL:
        undone = entry.reverses.kind if entry.reverses else None
        if undone == Kind.PAYMENT:
            return None, entry.amount
        if undone == Kind.REFUND:
            return entry.amount, None
        return None, None
    if kind == Kind.PAYMENT:
        return entry.amount, None
    if kind == Kind.REFUND:
        return None, entry.amount
    return None, None


def _fixture(entry: LedgerEntry) -> str:
    fixture = entry.fixture
    if fixture is None:
        return ""
    kick_off = timezone.localtime(fixture.kick_off)
    return f"{fixture.home_team.name} v {fixture.away_team.name}, {kick_off:%Y-%m-%d %H:%M}"


def row(entry: LedgerEntry) -> dict:
    """One entry as the export's columns. Money stays Decimal; dates stay dates."""
    recorded = timezone.localtime(entry.created_at)
    team = entry.team
    fixture = entry.fixture
    season = team.season if team else fixture.season if fixture else None
    division = team.division if team else fixture.division if fixture else None
    money_in, money_out = money_in_out(entry)
    return {
        "Entry": entry.pk,
        "Date": recorded.date(),
        "Time": f"{recorded:%H:%M}",
        "Kind": entry.get_kind_display(),
        "Category": entry.get_category_display(),
        "Method": entry.get_method_display() if entry.method else "",
        "Amount (R)": entry.amount,
        "Money in (R)": money_in,
        "Money out (R)": money_out,
        "Account": entry.get_subject_display(),
        "Team": safe_text(team.name) if team else "",
        "Player or Booking User": safe_text(entry.user.name) if entry.user else "",
        "Season": safe_text(season.name) if season else "",
        "Night": division.night.get_weekday_display() if division else "",
        "Division": safe_text(division.name) if division else "",
        "Fixture": safe_text(_fixture(entry)),
        "Booking": entry.booking_id,
        "Corrects entry": entry.reverses_id,
        "Credit from entry": entry.source_id,
        "Reason": safe_text(entry.reason),
        "Recorded by": safe_text(entry.recorded_by.name) if entry.recorded_by else "League App",
    }


def _csv_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value)


def to_csv(club: Club, period: Period) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(COLUMNS)
    for entry in entries(club, period).iterator():
        values = row(entry)
        writer.writerow([_csv_value(values[column]) for column in COLUMNS])
    # A BOM so Excel opens the UTF-8 (names like "Zoë") correctly.
    return out.getvalue().encode("utf-8-sig")


def export(club: Club, period: Period, file_type: str) -> tuple[bytes, str, str]:
    """The file's bytes, content type and name."""
    if file_type == "csv":
        return to_csv(club, period), "text/csv; charset=utf-8", f"{period.name}.csv"
    return to_xlsx(club, period), XLSX, f"{period.name}.xlsx"


def to_xlsx(club: Club, period: Period) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ledger"
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    count = 0
    for entry in entries(club, period).iterator():
        values = row(entry)
        sheet.append([values[column] for column in COLUMNS])
        count += 1
        for cell, column in zip(sheet[sheet.max_row], COLUMNS, strict=True):
            if column in MONEY_COLUMNS:
                cell.number_format = "#,##0.00"
            elif column == "Date":
                cell.number_format = "yyyy-mm-dd"
    for index, column in enumerate(COLUMNS, start=1):
        letter = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[letter].width = WIDTHS.get(column, max(12, len(column) + 2))

    about = workbook.create_sheet("About")
    exported = timezone.localtime()
    for line in [
        ("Club", safe_text(club.name)),
        ("From", period.starts_on),
        ("To", period.ends_on),
        ("Exported at", f"{exported:%Y-%m-%d %H:%M} (South African time)"),
        ("Entries", count),
        ("Currency", "South African rand (ZAR)"),
        (),
        ("Column", "Meaning"),
        *ABOUT,
    ]:
        about.append(list(line))
    about["A1"].font = Font(bold=True)
    for cell in about[8]:
        cell.font = Font(bold=True)
    about["B2"].number_format = about["B3"].number_format = "yyyy-mm-dd"
    about.column_dimensions["A"].width = 28
    about.column_dimensions["B"].width = 90

    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
