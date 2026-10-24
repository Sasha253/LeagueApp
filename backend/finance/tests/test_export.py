"""The accountant export (FR-51): a Club's ledger entries for any period, as CSV or Excel,
for its Admins only."""

import csv
import datetime as dt
import io
from decimal import Decimal

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import ClubRole
from finance import export
from ledger import services as ledger
from ledger.models import LedgerEntry

SAST = dt.timezone(dt.timedelta(hours=2))


def url(club):
    return reverse("api-v1:club-finance-export", kwargs={"club_id": club.pk})


def today() -> str:
    return timezone.localdate().isoformat()


def period(starts_on=None, ends_on=None, file_type=None):
    query = {"starts_on": starts_on or today(), "ends_on": ends_on or today()}
    if file_type:
        query["file_type"] = file_type
    return query


def csv_rows(response) -> list[dict]:
    text = response.content.decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(text)))


@pytest.mark.django_db
def test_csv_has_every_entry_of_the_club_oldest_first(books, admin_of, club_a):
    answer = admin_of(club_a).get(url(club_a), period(file_type="csv"))

    assert answer.status_code == 200
    assert answer["Content-Type"] == "text/csv; charset=utf-8"
    assert answer["Content-Disposition"] == (
        f'attachment; filename="ledger-{today()}-to-{today()}.csv"'
    )
    assert answer.content.startswith("﻿".encode())
    rows = csv_rows(answer)
    assert list(rows[0].keys()) == export.COLUMNS
    expected = LedgerEntry.objects.filter(club=club_a).order_by("created_at", "pk")
    assert [int(r["Entry"]) for r in rows] == [e.pk for e in expected]
    # Club B's R 900 never shows.
    assert all(r["Team"] != "B Team" for r in rows)
    assert "900.00" not in {r["Amount (R)"] for r in rows}


@pytest.mark.django_db
def test_csv_columns_for_a_charge_a_card_payment_and_a_correction(books, admin_of, club_a):
    rows = {int(r["Entry"]): r for r in csv_rows(
        admin_of(club_a).get(url(club_a), period(file_type="csv"))
    )}  # fmt: skip
    lions = books["lions"]

    charge = LedgerEntry.objects.get(team=lions, kind="charge")
    row = rows[charge.pk]
    local = timezone.localtime(charge.created_at)
    assert row["Date"] == local.date().isoformat()
    assert row["Time"] == f"{local:%H:%M}"
    assert row["Kind"] == "Charge"
    assert row["Category"] == "League"
    assert row["Method"] == ""
    assert row["Amount (R)"] == "350.00"
    assert row["Money in (R)"] == row["Money out (R)"] == ""
    assert row["Account"] == "Team"
    assert row["Team"] == "Red Lions"
    assert row["Season"] == "Spring"
    assert row["Night"] == "Tuesday"
    assert row["Division"] == "Division 2"
    assert row["Recorded by"] == "League App"

    card = LedgerEntry.objects.get(team=lions, kind="payment", method="card")
    row = rows[card.pk]
    assert row["Kind"] == "Payment"
    assert row["Method"] == "Card"
    assert row["Money in (R)"] == "100.00"
    assert row["Money out (R)"] == ""
    assert row["Account"] == "Player"
    assert row["Player or Booking User"] == "Sipho"

    returned = LedgerEntry.objects.get(source=card)
    assert rows[returned.pk]["Kind"] == "Credit"
    assert rows[returned.pk]["Credit from entry"] == str(card.pk)
    assert rows[returned.pk]["Money in (R)"] == rows[returned.pk]["Money out (R)"] == ""
    assert rows[returned.pk]["Reason"] == "Match cancelled"
    assert rows[returned.pk]["Recorded by"] == "Cash Admin"

    wrong = LedgerEntry.objects.get(team=books["boys"], kind="payment", method="cash")
    correction = LedgerEntry.objects.get(reverses=wrong)
    assert rows[wrong.pk]["Money in (R)"] == "50.00"
    row = rows[correction.pk]
    assert row["Kind"] == "Reversal"
    assert row["Corrects entry"] == str(wrong.pk)
    assert row["Money in (R)"] == ""
    assert row["Money out (R)"] == "50.00"
    assert row["Reason"] == "Wrong Team"


@pytest.mark.django_db
def test_money_in_and_out_is_only_card_and_cash(books, admin_of, club_a):
    rows = csv_rows(admin_of(club_a).get(url(club_a), period(file_type="csv")))

    def total(column):
        return sum((Decimal(r[column]) for r in rows if r[column]), Decimal("0"))

    # In: Lions 100 card + 250 cash, Old Boys 50 cash + 120 card, friendly 80 cash,
    # booking 600 card. Out: the Old Boys' 50 reversed. Credit spent moves no money.
    assert total("Money in (R)") == Decimal("1200.00")
    assert total("Money out (R)") == Decimal("50.00")
    spent = LedgerEntry.objects.get(team=books["boys"], kind="payment", method="credit")
    [row] = [r for r in rows if r["Entry"] == str(spent.pk)]
    assert row["Method"] == "Credit"
    assert row["Money in (R)"] == row["Money out (R)"] == ""


@pytest.mark.django_db
def test_friendly_and_booking_rows(books, admin_of, club_a):
    rows = csv_rows(admin_of(club_a).get(url(club_a), period(file_type="csv")))

    [friendly] = [r for r in rows if r["Category"] == "Friendly"]
    assert friendly["Player or Booking User"] == "Lerato"
    assert friendly["Season"] == "Spring"
    assert friendly["Division"] == "Division 2"
    kick_off = timezone.localtime(books["friendly"].kick_off)
    assert friendly["Fixture"] == f"Red Lions v Old Boys, {kick_off:%Y-%m-%d %H:%M}"

    bookings = [r for r in rows if r["Category"] == "Booking"]
    assert bookings
    assert {r["Account"] for r in bookings} == {"Booking User"}
    assert all(r["Booking"] and r["Season"] == r["Night"] == "" for r in bookings)


@pytest.mark.django_db
def test_period_is_south_african_days_both_inclusive(
    admin_of, club_a, make_team, people, time_machine
):  # fmt: skip
    team = make_team("Red Lions", people["thandi"])

    def paid_at(moment, amount):
        time_machine.move_to(moment, tick=False)
        return ledger.record_payment(
            team=team, amount=Decimal(amount), method="cash", recorded_by=None
        )

    # 23:59 on 31 Oct in South Africa is 21:59 UTC; 00:00 on 1 Nov is 22:00 UTC on 31 Oct.
    before = paid_at(dt.datetime(2026, 9, 30, 23, 59, tzinfo=SAST), "1")
    first = paid_at(dt.datetime(2026, 10, 1, 0, 0, tzinfo=SAST), "2")
    last = paid_at(dt.datetime(2026, 10, 31, 23, 59, tzinfo=SAST), "3")
    after = paid_at(dt.datetime(2026, 11, 1, 0, 0, tzinfo=SAST), "4")

    rows = csv_rows(admin_of(club_a).get(url(club_a), period("2026-10-01", "2026-10-31", "csv")))

    ids = [int(r["Entry"]) for r in rows]
    assert ids == [first.pk, last.pk]
    assert before.pk not in ids and after.pk not in ids
    assert rows[0]["Date"] == "2026-10-01" and rows[0]["Time"] == "00:00"
    assert rows[1]["Date"] == "2026-10-31" and rows[1]["Time"] == "23:59"


@pytest.mark.django_db
def test_an_empty_period_is_just_the_header(admin_of, club_a):
    answer = admin_of(club_a).get(url(club_a), period("2020-01-01", "2020-01-31", "csv"))

    assert answer.status_code == 200
    assert answer.content.decode("utf-8-sig").strip() == ",".join(export.COLUMNS)


@pytest.mark.django_db
def test_excel_has_the_ledger_and_an_about_sheet(books, admin_of, club_a):
    answer = admin_of(club_a).get(url(club_a), period())

    assert answer.status_code == 200
    assert answer["Content-Type"] == export.XLSX
    assert answer["Content-Disposition"].endswith(f'ledger-{today()}-to-{today()}.xlsx"')
    workbook = load_workbook(io.BytesIO(answer.content))
    assert workbook.sheetnames == ["Ledger", "About"]
    sheet = workbook["Ledger"]
    header = [cell.value for cell in sheet[1]]
    assert header == export.COLUMNS
    assert sheet.freeze_panes == "A2"
    rows = [dict(zip(header, (c.value for c in r), strict=True)) for r in sheet.iter_rows(2)]
    assert len(rows) == LedgerEntry.objects.filter(club=club_a).count()
    charge = LedgerEntry.objects.get(team=books["lions"], kind="charge")
    [row] = [r for r in rows if r["Entry"] == charge.pk]
    # Money is a number, the date a real date.
    assert row["Amount (R)"] == 350
    assert row["Money in (R)"] is None
    assert row["Date"].date() == timezone.localtime(charge.created_at).date()
    amount_cell = sheet.cell(row=rows.index(row) + 2, column=header.index("Amount (R)") + 1)
    assert amount_cell.number_format == "#,##0.00"

    about = {r[0].value: r[1].value for r in workbook["About"].iter_rows() if r[0].value}
    assert about["Club"] == "Club A"
    assert about["Entries"] == len(rows)
    assert about["Currency"] == "South African rand (ZAR)"


@pytest.mark.django_db
def test_text_that_looks_like_a_formula_is_escaped(admin_of, club_a, make_team, people):
    team = make_team('=HYPERLINK("http://x")', people["thandi"])
    admin = admin_of(club_a)
    ledger.issue_credit(
        club=club_a, amount=Decimal("5"), reason="+cmd|' /C calc'!A0", recorded_by=admin.user,
        team=team,
    )  # fmt: skip

    [row] = csv_rows(admin.get(url(club_a), period(file_type="csv")))
    assert row["Team"] == '\'=HYPERLINK("http://x")'
    assert row["Reason"] == "'+cmd|' /C calc'!A0"

    workbook = load_workbook(io.BytesIO(admin.get(url(club_a), period()).content))
    sheet = workbook["Ledger"]
    header = [cell.value for cell in sheet[1]]
    team_cell = sheet.cell(row=2, column=header.index("Team") + 1)
    assert team_cell.data_type == "s"
    assert team_cell.value == '\'=HYPERLINK("http://x")'


# The period


@pytest.mark.django_db
@pytest.mark.parametrize(
    "query, field",
    [
        ({"ends_on": "2026-10-31"}, "starts_on"),
        ({"starts_on": "2026-10-01"}, "ends_on"),
        ({"starts_on": "1 Oct", "ends_on": "2026-10-31"}, "starts_on"),
        ({"starts_on": "2026-10-31", "ends_on": "2026-10-01"}, "ends_on"),
        ({"starts_on": "2026-10-01", "ends_on": "2026-10-31", "file_type": "pdf"}, "file_type"),
    ],
)
def test_a_bad_period_is_refused(admin_of, club_a, query, field):
    answer = admin_of(club_a).get(url(club_a), query)

    assert answer.status_code == 400
    assert field in answer.json()


# Who may export


@pytest.mark.django_db
def test_only_the_clubs_admins_may_export(books, club_a, club_b, admin_of, make_user, give_role):
    from rest_framework.test import APIClient

    referee = make_user("export-referee@example.com")
    give_role(referee, club_a, ClubRole.Role.REFEREE)
    captain = books["lions"].captain
    for user in (referee, captain):
        client = APIClient()
        client.force_authenticate(user)
        assert client.get(url(club_a), period()).status_code == 403

    assert APIClient().get(url(club_a), period()).status_code == 401
    # Club B's Admin can't export Club A's ledger.
    assert admin_of(club_b).get(url(club_a), period()).status_code == 403


@pytest.mark.django_db
def test_an_unknown_club_is_not_found(admin_of, club_a):
    from rest_framework.test import APIClient

    client = APIClient()
    client.force_authenticate(admin_of(club_a).user)
    assert client.get("/api/v1/clubs/999999/finance/export/", period()).status_code in (403, 404)
