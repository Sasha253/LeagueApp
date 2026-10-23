"""Importing Nights and Divisions from a spreadsheet (FR-39)."""

import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from accounts.models import ClubRole
from clubs.models import Division, Night
from core import tables

HEADER = ["Night", "Division", "Min players", "League fee"]


def url(club, step):
    return reverse(f"api-v1:club-import-nights-{step}", kwargs={"club_id": club.pk})


def csv_file(*rows, header=HEADER, name="nights.csv"):
    lines = [",".join(header), *(",".join(f'"{cell}"' for cell in row) for row in rows)]
    return SimpleUploadedFile(name, "\n".join(lines).encode(), content_type="text/csv")


def xlsx_file(*rows, header=HEADER):
    workbook = Workbook()
    workbook.active.append(header)
    for row in rows:
        workbook.active.append(list(row))
    out = io.BytesIO()
    workbook.save(out)
    return SimpleUploadedFile("nights.xlsx", out.getvalue())


def upload(client, club, step, file):
    return client.post(url(club, step), {"file": file}, format="multipart")


def divisions(club):
    return sorted(
        (d.night.weekday, d.name, d.min_players, d.league_fee)
        for d in Division.objects.filter(night__club=club).select_related("night")
    )


# Reading the file


@pytest.mark.django_db
def test_preview_says_what_would_happen_and_writes_nothing(admin_of, club_a, tuesday):
    Division.objects.create(night=tuesday, name="Division 1", min_players=5)
    file = csv_file(
        ("Tuesday", "Division 1", "6", ""),
        ("Tuesday", "division 1 ", "", ""),
        ("Thu", "Division 1", "", "R 1,200.50"),
    )

    answer = upload(admin_of(club_a), club_a, "preview", file)

    assert answer.status_code == 200
    body = answer.json()
    assert body["committed"] is False
    assert body["summary"] == {
        "rows": 3,
        "nights_created": 1,
        "divisions_created": 1,
        "divisions_updated": 1,
        "unchanged": 0,
        "errors": 1,
    }
    update, duplicate, create = body["rows"]
    assert update == {
        "row": 2,
        "night": 1,
        "division": "Division 1",
        "min_players": 6,
        "league_fee": None,
        "action": "update",
        "new_night": False,
        "changes": ["min_players"],
        "errors": [],
    }
    assert duplicate["errors"] == ["This Division is already on row 2."]
    assert create["action"] == "create"
    assert create["new_night"] is True
    assert create["min_players"] == 5  # The club's minimum.
    assert create["league_fee"] == "1200.50"
    assert divisions(club_a) == [(1, "Division 1", 5, None)]
    assert club_a.nights.count() == 1


@pytest.mark.django_db
def test_each_bad_cell_is_reported_on_its_row(admin_of, club_a):
    file = csv_file(
        ("Someday", "Division 1", "", ""),
        ("", "", "", ""),  # Blank rows are skipped.
        ("Tuesday", "", "0", "350,50"),
        ("Tu", "D" * 61, "6.5", "-10"),
    )

    rows = upload(admin_of(club_a), club_a, "preview", file).json()["rows"]

    assert [row["row"] for row in rows] == [2, 4, 5]
    assert rows[0]["errors"] == ['Night "Someday" isn\'t a weekday.']
    assert rows[1]["errors"] == [
        "Division is empty.",
        "Min players must be a whole number of 1 or more.",
        "League fee must be an amount in rand, like 350 or 1,200.50.",
    ]
    assert rows[2]["errors"] == [
        'Night "Tu" isn\'t a weekday.',
        "Division is longer than 60 characters.",
        "Min players must be a whole number of 1 or more.",
        "League fee must be an amount in rand, like 350 or 1,200.50.",
    ]
    assert all(row["action"] == "error" for row in rows)


@pytest.mark.parametrize(
    ("text", "weekday"),
    [("Monday", 0), ("tue", 1), ("TUES", 1), ("Wednesday Night", 2), ("thurs", 3), ("Sun", 6)],
)
def test_weekdays_are_read_in_any_common_form(text, weekday):
    from clubs.imports import parse_weekday

    assert parse_weekday(text) == weekday


@pytest.mark.parametrize(
    ("text", "fee"),
    [
        ("350", "350.00"),
        ("R 350", "350.00"),
        ("R\u00a0350", "350.00"),
        ("R1 200", "1200.00"),
        ("1,200.5", "1200.50"),
    ],
)
def test_fees_are_read_as_rand(text, fee):
    from clubs.imports import parse_fee

    assert parse_fee(text) == Decimal(fee)


@pytest.mark.django_db
def test_xlsx_is_read_like_csv_and_headers_match_loosely(admin_of, club_a):
    file = xlsx_file(
        ("Tuesday", "Division 1", 6, 350),
        ("Friday", "Division 2", None, 400.5),
        header=["  NIGHT ", "Division name", "Minimum players", "League fee per Team", "Notes"],
    )

    rows = upload(admin_of(club_a), club_a, "preview", file).json()["rows"]

    assert [(r["night"], r["division"], r["min_players"], r["league_fee"]) for r in rows] == [
        (1, "Division 1", 6, "350.00"),
        (4, "Division 2", 5, "400.50"),
    ]


@pytest.mark.django_db
@pytest.mark.parametrize(
    ("file", "message"),
    [
        (csv_file(header=["Night", "Min players"]), 'The file has no "Division" column.'),
        (SimpleUploadedFile("nights.txt", b"Night,Division"), "Upload an .xlsx or .csv file."),
        (SimpleUploadedFile("nights.xlsx", b"not a workbook"), "can't be read as an Excel"),
        (SimpleUploadedFile("nights.csv", b",,\n\n"), "The file is empty."),
    ],
)
def test_a_file_that_cant_be_used_is_refused_whole(admin_of, club_a, file, message):
    answer = upload(admin_of(club_a), club_a, "preview", file)

    assert answer.status_code == 400
    assert message in answer.json()["file"][0]


@pytest.mark.django_db
def test_files_are_limited_in_size_and_rows(admin_of, club_a, monkeypatch):
    client = admin_of(club_a)
    monkeypatch.setattr(tables, "MAX_ROWS", 2)
    monkeypatch.setattr(tables, "MAX_BYTES", 200)

    too_long = csv_file(*[("Tuesday", f"Division {n}", "", "") for n in range(3)])
    assert "more than 2 rows" in upload(client, club_a, "preview", too_long).json()["file"][0]

    too_big = csv_file(("Tuesday", "D" * 300, "", ""))
    assert upload(client, club_a, "preview", too_big).json()["file"] == ["The file is over 5 MB."]


@pytest.mark.django_db
def test_semicolon_csv_from_excel_is_read(admin_of, club_a):
    file = SimpleUploadedFile("nights.csv", "﻿Night;Division\r\nTuesday;Division 1\r\n".encode())

    rows = upload(admin_of(club_a), club_a, "preview", file).json()["rows"]

    assert rows[0]["division"] == "Division 1"
    assert rows[0]["errors"] == []


# Saving it


@pytest.mark.django_db
def test_commit_saves_every_row_and_re_importing_changes_nothing(admin_of, club_a):
    client = admin_of(club_a)
    rows = [("Tuesday", "Division 1", "6", "350"), ("Thursday", "Division 1", "", "")]

    first = upload(client, club_a, "commit", csv_file(*rows))

    assert first.status_code == 200
    assert first.json()["committed"] is True
    assert divisions(club_a) == [
        (1, "Division 1", 6, Decimal("350.00")),
        (3, "Division 1", 5, None),
    ]

    again = upload(client, club_a, "commit", csv_file(*rows))

    assert again.json()["summary"]["unchanged"] == 2
    assert club_a.nights.count() == 2
    assert Division.objects.filter(night__club=club_a).count() == 2


@pytest.mark.django_db
def test_re_import_updates_values_keeps_blanks_and_leaves_other_divisions(
    admin_of, club_a, tuesday
):
    Division.objects.create(night=tuesday, name="Division 1", min_players=6, league_fee=350)
    Division.objects.create(night=tuesday, name="Division 2", min_players=7)

    upload(admin_of(club_a), club_a, "commit", csv_file(("TUESDAY", "DIVISION 1", "", "400")))

    assert divisions(club_a) == [
        (1, "Division 1", 6, Decimal("400.00")),  # Name as it was; blank min kept.
        (1, "Division 2", 7, None),  # Not in the file: left alone.
    ]


@pytest.mark.django_db
def test_one_bad_row_and_nothing_is_saved(admin_of, club_a):
    file = csv_file(("Tuesday", "Division 1", "", ""), ("Tuesday", "", "", ""))

    answer = upload(admin_of(club_a), club_a, "commit", file)

    assert answer.status_code == 400
    assert answer.json()["committed"] is False
    assert answer.json()["rows"][1]["errors"] == ["Division is empty."]
    assert Night.objects.filter(club=club_a).count() == 0


@pytest.mark.django_db
def test_commit_is_all_or_nothing_when_a_write_fails(admin_of, club_a, monkeypatch):
    real_create = Division.objects.create
    calls = []

    def fail_second(**fields):
        calls.append(fields)
        if len(calls) == 2:
            raise RuntimeError("The database went away")
        return real_create(**fields)

    monkeypatch.setattr(Division.objects, "create", fail_second)
    file = csv_file(("Tuesday", "Division 1", "", ""), ("Tuesday", "Division 2", "", ""))

    with pytest.raises(RuntimeError):
        upload(admin_of(club_a), club_a, "commit", file)

    assert Night.objects.filter(club=club_a).count() == 0
    assert Division.objects.filter(night__club=club_a).count() == 0


# Who may import


@pytest.mark.django_db
@pytest.mark.parametrize("step", ["preview", "commit"])
def test_only_this_clubs_admins_may_import(admin_of, club_a, club_b, make_user, give_role, step):
    from rest_framework.test import APIClient

    referee = make_user("ref@example.com")
    give_role(referee, club_a, ClubRole.Role.REFEREE)
    as_referee = APIClient()
    as_referee.force_authenticate(referee)
    file = lambda: csv_file(("Tuesday", "Division 1", "", ""))  # noqa: E731

    assert upload(admin_of(club_b), club_a, step, file()).status_code == 403
    assert upload(as_referee, club_a, step, file()).status_code == 403
    assert upload(APIClient(), club_a, step, file()).status_code == 401
    assert Night.objects.count() == 0


@pytest.mark.django_db
def test_an_import_only_touches_its_own_club(admin_of, club_a, club_b):
    other = Night.objects.create(club=club_b, weekday=Night.Weekday.TUESDAY)
    Division.objects.create(night=other, name="Division 1", min_players=9)

    answer = upload(admin_of(club_a), club_a, "commit", csv_file(("Tuesday", "Division 1", "", "")))

    assert answer.json()["rows"][0]["action"] == "create"
    assert divisions(club_b) == [(1, "Division 1", 9, None)]
    assert divisions(club_a) == [(1, "Division 1", 5, None)]


# The template


@pytest.mark.django_db
def test_the_template_downloads_as_xlsx_or_csv_and_imports_cleanly(admin_of, club_a, club_b):
    client = admin_of(club_a)

    xlsx = client.get(url(club_a, "template"))
    assert xlsx.status_code == 200
    assert xlsx["Content-Disposition"] == 'attachment; filename="nights-and-divisions.xlsx"'
    sheet = load_workbook(io.BytesIO(xlsx.content)).worksheets[0]
    assert [cell.value for cell in sheet[1]] == HEADER

    csv = client.get(url(club_a, "template"), {"file_type": "csv"})
    assert csv["Content-Type"] == "text/csv"
    assert csv.content.decode("utf-8-sig").splitlines()[0] == ",".join(HEADER)

    template = SimpleUploadedFile("template.xlsx", xlsx.content)
    assert upload(client, club_a, "preview", template).json()["summary"]["errors"] == 0
    assert admin_of(club_b).get(url(club_a, "template")).status_code == 403
