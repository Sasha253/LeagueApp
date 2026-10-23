"""Importing a Club's Nights and Divisions from the Admin's spreadsheet (FR-39).

One row per Division: Night (a weekday), Division, and optionally Min players and League
fee. A Division is matched on its Night and its name, ignoring case and outer spaces, so
importing the same file twice changes nothing (docs/DECISIONS.md):

- a new Division is created, and its Night too if the Club doesn't have it yet;
- an existing Division takes the file's Min players and League fee; a blank cell keeps the
  value it has;
- Divisions that aren't in the file are left alone.

plan() checks every row and writes nothing. apply() writes a plan with no errors.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.db import models
from openpyxl import Workbook
from openpyxl.styles import Font

from core.tables import Row

from .models import Club, Division, Night

COLUMNS = {
    "night": ("Night", "Weekday", "Day"),
    "division": ("Division", "Division name"),
    "min_players": ("Min players", "Minimum players"),
    "league_fee": ("League fee", "League fee per Team", "Fee"),
}
REQUIRED = {"night", "division"}

TEMPLATE_ROWS = [
    ("Tuesday", "Division 1", 6, 350),
    ("Tuesday", "Division 2", None, None),
    ("Thursday", "Division 1", None, 400),
]
TEMPLATE_HELP = [
    ("Column", "What to put in it"),
    ("Night", "The weekday: Monday to Sunday, or Mon to Sun. Required."),
    ("Division", "The Division's name, e.g. Division 1. Required."),
    ("Min players", "Fewer confirmed Players and the Team is out. Blank: the club's minimum."),
    ("League fee", "Fee per Team in rand, e.g. 350. Blank: the club's league fee."),
    ("", "Re-importing the same file creates nothing twice."),
    ("", "A blank cell keeps what a Division already has."),
    ("", "Divisions missing from the file are left as they are."),
]

NAME_MAX = Division._meta.get_field("name").max_length
SMALL_INT_MAX = 32767
FEE_MAX = Decimal("99999999.99")  # The ledger's money fields: 10 digits, 2 after the point.
FEE_FORMAT = re.compile(r"^(\d{1,3}(,\d{3})+|\d+)(\.\d{1,2})?$")
WEEKDAYS = {label[:3].lower(): value for value, label in Night.Weekday.choices}


class Action(models.TextChoices):
    CREATE = "create", "New Division"
    UPDATE = "update", "Division changes"
    UNCHANGED = "unchanged", "No change"
    ERROR = "error", "Has errors"


@dataclass
class PlannedRow:
    row: int
    night: int | None = None  # The weekday, 0 for Monday.
    division: str = ""
    min_players: int | None = None  # What the Division will have, after the import.
    league_fee: Decimal | None = None  # Its own fee; None means the club's.
    action: str = Action.ERROR
    new_night: bool = False
    changes: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    existing: Division | None = None


@dataclass
class Plan:
    rows: list[PlannedRow]

    @property
    def has_errors(self) -> bool:
        return any(row.errors for row in self.rows)

    @property
    def summary(self) -> dict[str, int]:
        def count(action):
            return sum(1 for row in self.rows if row.action == action)

        return {
            "rows": len(self.rows),
            "nights_created": len({row.night for row in self.rows if row.new_night}),
            "divisions_created": count(Action.CREATE),
            "divisions_updated": count(Action.UPDATE),
            "unchanged": count(Action.UNCHANGED),
            "errors": count(Action.ERROR),
        }


def parse_weekday(text: str) -> int | None:
    word = re.sub(r"\s+night$", "", text.strip().lower())
    if len(word) < 3:
        return None
    value = WEEKDAYS.get(word[:3])
    if value is None or not Night.Weekday(value).label.lower().startswith(word):
        return None
    return value


def parse_min_players(text: str) -> int | None:
    if not re.fullmatch(r"\d+", text):
        raise ValueError("Min players must be a whole number of 1 or more.")
    value = int(text)
    if not 1 <= value <= SMALL_INT_MAX:
        raise ValueError("Min players must be a whole number of 1 or more.")
    return value


def parse_fee(text: str) -> Decimal:
    plain = re.sub(r"\s", "", text)  # Unicode \s takes no-break spaces too.
    plain = re.sub(r"^[Rr]", "", plain)
    if not FEE_FORMAT.fullmatch(plain):
        raise ValueError("League fee must be an amount in rand, like 350 or 1,200.50.")
    try:
        value = Decimal(plain.replace(",", ""))
    except InvalidOperation as error:
        raise ValueError("League fee must be an amount in rand, like 350 or 1,200.50.") from error
    if value > FEE_MAX:
        raise ValueError("League fee is too large.")
    return value.quantize(Decimal("0.01"))


def plan(club: Club, rows: list[Row]) -> Plan:
    nights = {night.weekday: night for night in club.nights.all()}
    divisions = {
        (division.night.weekday, division.name.strip().casefold()): division
        for division in Division.objects.filter(night__club=club).select_related("night")
    }
    default_min = club.settings.division_min_players
    seen: dict[tuple[int, str], int] = {}
    planned = []

    for row in rows:
        result = PlannedRow(row=row.number)
        planned.append(result)
        cells = row.cells

        night_text = cells.get("night", "")
        result.night = parse_weekday(night_text)
        if not night_text:
            result.errors.append("Night is empty.")
        elif result.night is None:
            result.errors.append(f'Night "{night_text}" isn\'t a weekday.')

        result.division = cells.get("division", "")
        if not result.division:
            result.errors.append("Division is empty.")
        elif len(result.division) > NAME_MAX:
            result.errors.append(f"Division is longer than {NAME_MAX} characters.")

        min_players = fee = None
        if cells.get("min_players"):
            try:
                min_players = parse_min_players(cells["min_players"])
            except ValueError as error:
                result.errors.append(str(error))
        if cells.get("league_fee"):
            try:
                fee = parse_fee(cells["league_fee"])
            except ValueError as error:
                result.errors.append(str(error))

        if result.errors:
            continue
        key = (result.night, result.division.casefold())
        if key in seen:
            result.errors.append(f"This Division is already on row {seen[key]}.")
            continue
        seen[key] = row.number

        existing = divisions.get(key)
        result.existing = existing
        result.new_night = result.night not in nights
        if existing is None:
            result.action = Action.CREATE
            result.min_players = min_players if min_players is not None else default_min
            result.league_fee = fee
            continue
        result.division = existing.name
        result.min_players = min_players if min_players is not None else existing.min_players
        result.league_fee = fee if fee is not None else existing.league_fee
        if result.min_players != existing.min_players:
            result.changes.append("min_players")
        if result.league_fee != existing.league_fee:
            result.changes.append("league_fee")
        result.action = Action.UPDATE if result.changes else Action.UNCHANGED

    return Plan(planned)


def apply(club: Club, import_plan: Plan) -> None:
    """Writes a plan with no errors. Call it in a transaction."""
    if import_plan.has_errors:
        raise ValueError("A plan with errors can't be applied.")
    nights = {night.weekday: night for night in club.nights.all()}
    for row in import_plan.rows:
        if row.action == Action.CREATE:
            night = nights.get(row.night)
            if night is None:
                night = nights[row.night] = Night.objects.create(club=club, weekday=row.night)
            Division.objects.create(
                night=night,
                name=row.division,
                min_players=row.min_players,
                league_fee=row.league_fee,
            )
        elif row.action == Action.UPDATE:
            division = row.existing
            division.min_players = row.min_players
            division.league_fee = row.league_fee
            division.save(update_fields=["min_players", "league_fee"])


def template_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Nights and Divisions"
    sheet.append([headers[0] for headers in COLUMNS.values()])
    for row in TEMPLATE_ROWS:
        sheet.append(list(row))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for letter, width in zip("ABCD", (14, 24, 14, 14), strict=True):
        sheet.column_dimensions[letter].width = width

    help_sheet = workbook.create_sheet("How to fill it in")
    for row in TEMPLATE_HELP:
        help_sheet.append(list(row))
    help_sheet["A1"].font = help_sheet["B1"].font = Font(bold=True)
    help_sheet.column_dimensions["A"].width = 14
    help_sheet.column_dimensions["B"].width = 80

    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def template_csv() -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow([headers[0] for headers in COLUMNS.values()])
    for row in TEMPLATE_ROWS:
        writer.writerow(["" if value is None else value for value in row])
    return out.getvalue().encode("utf-8-sig")
